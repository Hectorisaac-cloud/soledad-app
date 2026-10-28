"""Carga del historial diario de julio y agosto-septiembre a la bitácora.

Lee el archivo de campo, valida cada ubicación contra el catálogo maestro y
separa, en los comentarios, el consumo real de las notas preventivas.
"""

from __future__ import annotations

import hashlib
import os
import re
import unicodedata
from datetime import date, datetime, timedelta

from openpyxl import load_workbook

RAIZ = os.path.dirname(os.path.abspath(__file__))
CANDIDATOS = (
    os.path.join(RAIZ, "assets", "historial_pjp4.xlsx"),
    os.path.join(RAIZ, "assets", "historial_pjp4.xlsx.xlsx"),
)

HOJAS_OBJETIVO = (
    ("JULIO", lambda clave: clave == "JULIO" or clave.startswith("JULIO ")),
    ("AGOSTO-SEPTIEMBRE", lambda clave: "AGOSTO" in clave and "SEPTIEMBRE" in clave),
)

_PALABRAS = {
    "UN": 1,
    "UNA": 1,
    "UNO": 1,
    "DOS": 2,
    "TRES": 3,
    "CUATRO": 4,
    "CINCO": 5,
    "SEIS": 6,
    "SIETE": 7,
    "OCHO": 8,
    "NUEVE": 9,
    "DIEZ": 10,
    "ONCE": 11,
    "DOCE": 12,
    "QUINCE": 15,
    "VEINTE": 20,
    "TREINTA": 30,
}
_MODELOS = {
    "112": "ciento doce",
    "118": "ciento dieciocho",
    "125": "ciento veinticinco",
    "144": "ciento cuarenta y cuatro",
    "270": "doscientos setenta",
}
_DIGITOS = ("cero", "uno", "dos", "tres", "cuatro", "cinco", "seis", "siete", "ocho", "nueve")
_RE_VOLUMEN = re.compile(r"(\d+(?:[.,]\d+)?)\s*(?:LITROS?|LTS?|LTROS?|L\b)")
_RE_VOLUMEN_PALABRA = re.compile(
    r"\b(" + "|".join(_PALABRAS) + r")\s+(?:LITROS?|LTS?|LTROS?)\b"
)
_RE_KG = re.compile(r"(\d+(?:[.,]\d+)?)\s*K(?:G|GS|ILOS?)\b")
_RE_MODELO = re.compile(r"\bC\s*-?\s*(\d{2,3})\b")
_RE_CONSUMO = re.compile(
    r"\b("
    r"SE INSTAL\w*|INSTALACION|"
    r"SE COLOC\w*|COLOCACION|"
    r"SE AGREG\w*|AGREGANDO|"
    r"SUMI+N\w*|"
    r"RE+MPLAZO\w*|RE+MPLAZ\w*|"
    r"CAMBIO DE|SE CAMBI\w*|"
    r"SE APLIC\w*"
    r")\b"
)
_RE_PREVENTIVA = re.compile(
    r"\b("
    r"AJUSTE|LIMPIEZA|INSPECCION|MEDICION|GROUND TESTER|OHMS?|"
    r"REQUIERE|REQUIEREN|PENDIENTE|OPERA CON"
    r")\b"
)


def ruta_archivo() -> str:
    for ruta in CANDIDATOS:
        if os.path.isfile(ruta):
            return ruta
    return ""


def _sin(texto) -> str:
    base = unicodedata.normalize("NFD", str(texto or ""))
    base = "".join(c for c in base if unicodedata.category(c) != "Mn")
    return " ".join(base.upper().split())


def _clave_hoja(nombre: str) -> str:
    return " ".join(_sin(nombre).replace("-", " ").replace("_", " ").split())


def _numero(valor) -> float:
    if isinstance(valor, bool) or valor is None:
        return 0.0
    if isinstance(valor, (int, float)):
        numero = float(valor)
        if numero != numero or numero <= 0:
            return 0.0
        return numero
    texto = str(valor).strip().replace(",", ".")
    if not texto or texto.lower() in {"nan", "none", "-", "nat"}:
        return 0.0
    try:
        numero = float(texto.replace("%", ""))
    except ValueError:
        return 0.0
    return numero if numero > 0 else 0.0


def _fecha(valor) -> date | None:
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    if isinstance(valor, (int, float)) and not isinstance(valor, bool):
        serial = float(valor)
        if 20000 < serial < 80000:
            return (datetime(1899, 12, 30) + timedelta(days=serial)).date()
    texto = str(valor or "").strip()
    if not texto or texto.lower() in {"nan", "none", "nat"}:
        return None
    for fmt in ("%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d", "%d/%m/%y", "%Y/%m/%d"):
        try:
            return datetime.strptime(texto[:10], fmt).date()
        except ValueError:
            continue
    return None


def _avance(valor) -> float:
    numero = _numero(valor)
    if numero <= 0:
        return 0.0
    if numero <= 1:
        return round(numero * 100, 2)
    return numero


def _volumen(clausula: str) -> float:
    hallado = _RE_VOLUMEN.search(clausula)
    if hallado:
        return float(hallado.group(1).replace(",", "."))
    palabra = _RE_VOLUMEN_PALABRA.search(clausula)
    if palabra:
        return float(_PALABRAS[palabra.group(1)])
    return 0.0


def _kilogramos(clausula: str) -> float:
    hallado = _RE_KG.search(clausula)
    if not hallado:
        return 0.0
    return float(hallado.group(1).replace(",", "."))


def _piezas(clausula: str, sustantivo: str) -> tuple[float, bool]:
    cifra = re.search(rf"(\d+(?:[.,]\d+)?)\s*(?:DE\s+)?(?:{sustantivo})", clausula)
    if cifra:
        return float(cifra.group(1).replace(",", ".")), True
    palabra = re.search(rf"\b({'|'.join(_PALABRAS)})\s+(?:DE\s+)?(?:{sustantivo})", clausula)
    if palabra:
        return float(_PALABRAS[palabra.group(1)]), True
    if re.search(rf"\b(?:{sustantivo})\b", clausula):
        return 1.0, False
    return 0.0, False


def _clausulas(texto: str) -> list[str]:
    plano = _sin(texto)
    if not plano:
        return []
    salida = []
    for frase in re.split(r"[.;]+", plano):
        for parte in frase.split(","):
            parte = " ".join(parte.split())
            if parte:
                salida.append(parte)
    return salida


def _es_preventiva(clausula: str) -> bool:
    return bool(_RE_PREVENTIVA.search(clausula))


def _es_consumo(clausula: str) -> bool:
    return bool(_RE_CONSUMO.search(clausula))


def _hereda_consumo(clausula: str) -> bool:
    if _es_preventiva(clausula) or _es_consumo(clausula):
        return False
    if re.search(r"\b(OPERA|QUEDA|PRESENTA|NIVEL|MOTOR|UNIDAD|NOTA|SE )\b", clausula):
        return False
    return True


def _origen(clausula: str) -> str:
    if "PJP4" in clausula:
        return "PJP4"
    if re.search(r"\bOPS\b", clausula):
        return "OPS"
    return ""


def _modelo_palabras(numero: str) -> str:
    if numero in _MODELOS:
        return _MODELOS[numero]
    return " ".join(_DIGITOS[int(digito)] for digito in numero)


def _nota_banda(clausula: str) -> str:
    if "RECUPERAD" in clausula:
        return "Recuperadas"
    if "ALTERNADOR" in clausula:
        return "Alternador"
    if "CORTA" in clausula:
        return "Cortas"
    return ""


def _sumar_bandas(eventos: list[dict]) -> list[dict]:
    if not eventos:
        return []
    explicitos = [evento for evento in eventos if evento["explicita"]]
    if not explicitos:
        elegido = dict(eventos[-1])
        for evento in eventos:
            if evento["nota"] and not elegido["nota"]:
                elegido["nota"] = evento["nota"]
            if evento["origen"] and not elegido["origen"]:
                elegido["origen"] = evento["origen"]
            if evento["modelo"] and not elegido["modelo"]:
                elegido["modelo"] = evento["modelo"]
        return [elegido]
    unidos: dict[str, dict] = {}
    for evento in explicitos:
        clave = evento["modelo"] or evento["nota"] or "banda"
        actual = unidos.get(clave)
        if actual is None:
            unidos[clave] = dict(evento)
            continue
        actual["cantidad"] += evento["cantidad"]
        if not actual["origen"] and evento["origen"]:
            actual["origen"] = evento["origen"]
    return list(unidos.values())


def extraer_insumos(comentario: str) -> dict:
    """Consumo real del comentario. Limpieza, ajuste y 'requiere' no cuentan."""
    insumos = {
        "aceite": 0.0,
        "hidraulico": 0.0,
        "refrigerante": 0.0,
        "grasa": 0.0,
        "filtros": 0.0,
        "bujias": 0.0,
        "bomba_agua": 0.0,
        "alternador": 0.0,
        "bandas": [],
    }
    verbos = False
    bandas: list[dict] = []
    bomba_implicita = False
    for clausula in _clausulas(comentario):
        if _es_preventiva(clausula):
            verbos = False
            continue
        if _es_consumo(clausula):
            verbos = True
        elif verbos and _hereda_consumo(clausula):
            pass
        else:
            verbos = False
            continue
        if "BOMBA DE AGUA" in clausula and not bomba_implicita:
            insumos["bomba_agua"] += 1.0
            bomba_implicita = True
        if re.search(r"\bBANDA", clausula):
            cantidad, explicita = _piezas(clausula, r"BANDAS?")
            if cantidad <= 0:
                cantidad, explicita = 1.0, False
            modelos = _RE_MODELO.findall(clausula)
            nota = _nota_banda(clausula)
            origen = _origen(clausula)
            if modelos:
                for modelo in modelos:
                    bandas.append(
                        {
                            "modelo": modelo,
                            "cantidad": cantidad if len(modelos) == 1 else 1.0,
                            "explicita": explicita if len(modelos) == 1 else False,
                            "origen": origen,
                            "nota": nota,
                        }
                    )
            else:
                bandas.append(
                    {
                        "modelo": "",
                        "cantidad": cantidad,
                        "explicita": explicita,
                        "origen": origen,
                        "nota": nota,
                    }
                )
        if re.search(r"ANTICONGELANTE|REFRIGERANTE", clausula) and "OK" not in clausula:
            insumos["refrigerante"] += _volumen(clausula)
        if "GRASA" in clausula and "RETIRO" not in clausula and "EXCEDENTE" not in clausula:
            insumos["grasa"] += _kilogramos(clausula)
        if "ACEITE HIDRAULICO" in clausula or ("HIDRAULICO" in clausula and "ACEITE" in clausula):
            insumos["hidraulico"] += _volumen(clausula)
        elif (
            "ACEITE" in clausula
            and "FILTRO" not in clausula
            and "INDICADOR" not in clausula
            and "PRESION" not in clausula
            and "EMULSION" not in clausula
            and "CONTAMIN" not in clausula
        ):
            insumos["aceite"] += _volumen(clausula)
        if re.search(
            r"(COLOCACION|INSTALACION|SE COLOCA|SE INSTALA)\s+(?:DE\s+)?(?:EL\s+|LA\s+|UN\s+|UNA\s+)?FILTRO",
            clausula,
        ):
            insumos["filtros"] += 1.0
        if "BUJIA" in clausula and not re.search(r"LIMPIEZA|CABLE", clausula):
            cantidad, _explicita = _piezas(clausula, r"BUJIAS?")
            insumos["bujias"] += cantidad or 1.0
        if "ALTERNADOR" in clausula and "BANDA" not in clausula and "DANAD" not in clausula:
            insumos["alternador"] += 1.0
    insumos["bandas"] = _sumar_bandas(bandas)
    return insumos


def _texto_bandas(bandas: list[dict]) -> tuple[float, str, str, str]:
    c144 = 0.0
    descripciones = []
    otros = []
    origenes = []
    for banda in bandas:
        cantidad = float(banda["cantidad"] or 0)
        if cantidad <= 0:
            continue
        modelo = str(banda.get("modelo") or "")
        nota = str(banda.get("nota") or "")
        origen = str(banda.get("origen") or "")
        if origen and origen not in origenes:
            origenes.append(origen)
        if modelo == "144":
            c144 += cantidad
            descripciones.append(f"C-144 x {cantidad:g}")
            continue
        if modelo:
            descripciones.append(f"C-{modelo} x {cantidad:g}")
            otros.append(f"Bandas C {_modelo_palabras(modelo)} x {cantidad:g}")
        else:
            etiqueta = nota or "Bandas"
            descripciones.append(f"{etiqueta} x {cantidad:g}")
            otros.append(f"Bandas {etiqueta.lower()} x {cantidad:g}" if nota else f"Bandas x {cantidad:g}")
    return c144, ", ".join(descripciones), ", ".join(origenes), ", ".join(otros)


def _id_registro(fecha: date, cuadrilla: str, ubicacion: str, descripcion: str, comentario: str) -> str:
    base = "|".join(
        [
            fecha.isoformat(),
            _sin(cuadrilla),
            _sin(ubicacion),
            _sin(descripcion),
            _sin(comentario),
        ]
    )
    digest = hashlib.sha1(base.encode("utf-8", errors="ignore")).hexdigest()[:10].upper()
    return f"HIST-{fecha.strftime('%Y%m%d')}-{digest}"


def _encabezado(hoja) -> tuple[int, dict[str, int]]:
    for indice, fila in enumerate(hoja.iter_rows(min_row=1, max_row=15, max_col=16, values_only=True), start=1):
        mapa: dict[str, int] = {}
        for columna, valor in enumerate(fila, start=1):
            clave = _sin(valor)
            if not clave:
                continue
            if "FECHA" in clave and "fecha" not in mapa:
                mapa["fecha"] = columna
            elif clave.startswith("CUADRILLA"):
                mapa["cuadrilla"] = columna
            elif "UBICACION" in clave:
                mapa["ubicacion"] = columna
            elif "DESCRIPCION" in clave:
                mapa["descripcion"] = columna
            elif "COMENTARIO" in clave:
                mapa["comentarios"] = columna
            elif "AVANCE" in clave:
                mapa["avance"] = columna
            elif "ACEITE" in clave and "PJP4" in clave:
                mapa["aceite_pjp4"] = columna
            elif "ACEITE" in clave and "OPS" in clave:
                mapa["aceite_ops"] = columna
        if {"fecha", "cuadrilla", "ubicacion", "comentarios"} <= set(mapa):
            return indice, mapa
    return 0, {}


def _elegir_hojas(libro) -> list[tuple[str, str]]:
    elegidas = []
    for etiqueta, coincide in HOJAS_OBJETIVO:
        candidatas = [nombre for nombre in libro.sheetnames if coincide(_clave_hoja(nombre))]
        diaria = ""
        for nombre in sorted(candidatas, key=lambda item: (0 if "2026" in item else 1, len(item))):
            fila, _mapa = _encabezado(libro[nombre])
            if fila:
                diaria = nombre
                break
        elegidas.append((etiqueta, diaria))
    return elegidas


def _celda(fila: tuple, columna: int):
    if columna <= 0 or columna > len(fila):
        return None
    return fila[columna - 1]


def _resumen_vacio(etiqueta: str, hoja: str) -> dict:
    return {
        "etiqueta": etiqueta,
        "hoja": hoja,
        "leidos": 0,
        "importados": 0,
        "repetidos": 0,
        "sin_fecha": 0,
        "sin_ubicacion": 0,
        "ubicacion_invalida": 0,
        "con_insumos": 0,
        "aceite_pjp4": 0.0,
        "aceite_ops": 0.0,
        "refrigerante": 0.0,
        "bandas": 0.0,
        "grasa": 0.0,
        "filtros": 0.0,
        "bujias": 0.0,
        "bombas": 0.0,
        "rechazos": {},
        "error": "",
    }


def _fila_bitacora(
    fecha: date,
    cuadrilla: str,
    ubicacion: str,
    descripcion: str,
    comentario: str,
    avance: float,
    aceite_pjp4: float,
    aceite_ops: float,
    insumos: dict,
) -> dict:
    c144, descripcion_bandas, origen_bandas, otros_bandas = _texto_bandas(insumos["bandas"])
    otros = [otros_bandas] if otros_bandas else []
    if aceite_ops > 0:
        otros.append(f"Aceite OPS x {aceite_ops:g}")
    if insumos["hidraulico"] > 0:
        otros.append(f"Aceite hidraulico x {insumos['hidraulico']:g}")
    total_bandas = sum(float(item["cantidad"] or 0) for item in insumos["bandas"])
    return {
        "FECHA": fecha,
        "CUADRILLA": cuadrilla,
        "PERSONAL": "",
        "UBICACIÓN": ubicacion,
        "EQUIPO/UBICACIÓN": ubicacion,
        "POZO": ubicacion,
        "SECTOR": "",
        "SECUENCIA": "",
        "DESCRIPCION DE LA ACTIVIDAD": descripcion,
        "DESCRIPCIÓN DE LA ACTIVIDAD": descripcion,
        "ACEITE SAE 50 PJP4": aceite_pjp4,
        "ACEITE (LTS)": aceite_pjp4,
        "ACEITE HIDRÁULICO (LTS)": insumos["hidraulico"],
        "BANDAS C144 PJP4": c144,
        "BANDAS CANTIDAD": c144,
        "BANDAS DESCRIPCIÓN": descripcion_bandas,
        "BANDAS PROPORCIONADO POR": origen_bandas,
        "BOMBA DE AGUA REPARADA PJP4": insumos["bomba_agua"],
        "GRASA OPS": insumos["grasa"],
        "GRASA (KG)": insumos["grasa"],
        "REFRIGERANTE OPS": insumos["refrigerante"],
        "REFRIGERANTE (LTS)": insumos["refrigerante"],
        "FILTROS (PZAS)": insumos["filtros"],
        "BUJÍAS (PZAS)": insumos["bujias"],
        "ALTERNADOR CANTIDAD": insumos["alternador"],
        "OTROS INSUMOS": ", ".join(parte for parte in otros if parte),
        "COMENTARIOS": comentario,
        "AVANCE %": avance,
        "ID_REGISTRO": _id_registro(fecha, cuadrilla, ubicacion, descripcion, comentario),
        "_bandas_total": total_bandas,
    }


def listar_ubicaciones(ruta: str) -> list[str]:
    """Ubicaciones distintas de las hojas diarias, en el orden en que aparecen."""
    if not ruta or not os.path.isfile(ruta):
        return []
    vistos: list[str] = []
    ya: set[str] = set()
    libro = load_workbook(ruta, read_only=True, data_only=True)
    try:
        for _etiqueta, nombre in _elegir_hojas(libro):
            if not nombre:
                continue
            hoja = libro[nombre]
            encabezado, mapa = _encabezado(hoja)
            columna = mapa.get("ubicacion", 0)
            if not encabezado or not columna:
                continue
            limite = min(int(hoja.max_row or encabezado), encabezado + 8000)
            vacios = 0
            for valores in hoja.iter_rows(
                min_row=encabezado + 1, max_row=limite, max_col=16, values_only=True
            ):
                texto = " ".join(str(_celda(valores, columna) or "").split())
                if texto.lower() in {"", "nan", "none"}:
                    vacios += 1
                    if vacios >= 25:
                        break
                    continue
                vacios = 0
                clave = _sin(texto)
                if clave in ya:
                    continue
                ya.add(clave)
                vistos.append(texto)
    finally:
        libro.close()
    return vistos


def procesar(ruta: str, resolver, ids_existentes: set[str] | None = None) -> dict:
    """Devuelve las filas nuevas y el resumen por hoja. No escribe la bitácora."""
    if not ruta or not os.path.isfile(ruta):
        return {"archivo": ruta or "", "hojas": [], "filas": [], "error": "No se encontró el archivo histórico."}
    conocidos = set(ids_existentes or ())
    vistos: set[str] = set()
    filas = []
    resumenes = []
    libro = load_workbook(ruta, read_only=True, data_only=True)
    try:
        for etiqueta, nombre in _elegir_hojas(libro):
            resumen = _resumen_vacio(etiqueta, nombre)
            if not nombre:
                resumen["error"] = f"La hoja de {etiqueta} no tiene el formato de bitácora diaria."
                resumenes.append(resumen)
                continue
            hoja = libro[nombre]
            encabezado, mapa = _encabezado(hoja)
            limite = min(int(hoja.max_row or encabezado), encabezado + 8000)
            vacios = 0
            for numero, valores in enumerate(
                hoja.iter_rows(min_row=encabezado + 1, max_row=limite, max_col=16, values_only=True),
                start=encabezado + 1,
            ):
                fecha = _fecha(_celda(valores, mapa.get("fecha", 0)))
                cuadrilla = " ".join(str(_celda(valores, mapa.get("cuadrilla", 0)) or "").split())
                ubicacion_bruta = " ".join(str(_celda(valores, mapa.get("ubicacion", 0)) or "").split())
                descripcion = " ".join(str(_celda(valores, mapa.get("descripcion", 0)) or "").split())
                comentario = " ".join(str(_celda(valores, mapa.get("comentarios", 0)) or "").split())
                if ubicacion_bruta.lower() in {"nan", "none"}:
                    ubicacion_bruta = ""
                if descripcion.lower() in {"nan", "none"}:
                    descripcion = ""
                if comentario.lower() in {"nan", "none"}:
                    comentario = ""
                if cuadrilla.lower() in {"nan", "none"}:
                    cuadrilla = ""
                aceite_col_pjp4 = _numero(_celda(valores, mapa.get("aceite_pjp4", 0)))
                aceite_col_ops = _numero(_celda(valores, mapa.get("aceite_ops", 0)))
                if not any((fecha, cuadrilla, ubicacion_bruta, descripcion, comentario, aceite_col_pjp4, aceite_col_ops)):
                    vacios += 1
                    if vacios >= 25:
                        break
                    continue
                vacios = 0
                resumen["leidos"] += 1
                if fecha is None:
                    resumen["sin_fecha"] += 1
                    continue
                if not ubicacion_bruta:
                    resumen["sin_ubicacion"] += 1
                    continue
                ubicacion = resolver(ubicacion_bruta) if resolver else ""
                if not ubicacion:
                    resumen["ubicacion_invalida"] += 1
                    resumen["rechazos"][ubicacion_bruta] = resumen["rechazos"].get(ubicacion_bruta, 0) + 1
                    continue
                insumos = extraer_insumos(comentario)
                aceite_pjp4 = aceite_col_pjp4
                if aceite_pjp4 <= 0 and aceite_col_ops <= 0 and insumos["aceite"] > 0:
                    aceite_pjp4 = insumos["aceite"]
                registro = _fila_bitacora(
                    fecha,
                    cuadrilla,
                    str(ubicacion).strip(),
                    descripcion,
                    comentario,
                    _avance(_celda(valores, mapa.get("avance", 0))),
                    aceite_pjp4,
                    aceite_col_ops,
                    insumos,
                )
                identificador = registro["ID_REGISTRO"]
                if identificador in conocidos or identificador in vistos:
                    resumen["repetidos"] += 1
                    continue
                vistos.add(identificador)
                resumen["importados"] += 1
                bandas_total = registro.pop("_bandas_total")
                if any(
                    (
                        aceite_pjp4,
                        aceite_col_ops,
                        insumos["refrigerante"],
                        insumos["grasa"],
                        insumos["hidraulico"],
                        insumos["filtros"],
                        insumos["bujias"],
                        insumos["bomba_agua"],
                        insumos["alternador"],
                        bandas_total,
                    )
                ):
                    resumen["con_insumos"] += 1
                resumen["aceite_pjp4"] += aceite_pjp4
                resumen["aceite_ops"] += aceite_col_ops
                resumen["refrigerante"] += insumos["refrigerante"]
                resumen["bandas"] += bandas_total
                resumen["grasa"] += insumos["grasa"]
                resumen["filtros"] += insumos["filtros"]
                resumen["bujias"] += insumos["bujias"]
                resumen["bombas"] += insumos["bomba_agua"]
                filas.append(registro)
            resumenes.append(resumen)
    finally:
        libro.close()
    return {"archivo": ruta, "hojas": resumenes, "filas": filas, "error": ""}
