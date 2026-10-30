"""Reporte fotográfico consolidado a partir del machote oficial."""

from __future__ import annotations

import io
import os
import re
import unicodedata
import warnings
import zipfile
from datetime import date, datetime

from openpyxl import load_workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor, TwoCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter
from openpyxl.utils.units import pixels_to_EMU
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet
from PIL import Image as PILImage
from PIL import ImageDraw, ImageOps
from xml.sax.saxutils import escape

from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.utils import ImageReader
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas
from reportlab.platypus import Paragraph, Table, TableStyle

import sitios_operativos

VERSION_DISENO_PDF = 27
PLANTILLA_XLSX = os.path.join("database", "machote_fotografico_template.xlsx")
FOTOS_POR_FILA = 3
FOTOS_POR_HOJA = 9
RANURAS = (
    (1, 10, 4, 23),
    (6, 10, 9, 23),
    (11, 10, 14, 23),
    (1, 25, 4, 38),
    (6, 25, 9, 38),
    (11, 25, 14, 38),
    (1, 40, 4, 53),
    (6, 40, 9, 53),
    (11, 40, 14, 53),
)
FIRMA_DERECHA = TwoCellAnchor(
    _from=AnchorMarker(col=9, colOff=243840, row=62, rowOff=38100),
    to=AnchorMarker(col=11, colOff=647700, row=66, rowOff=617220),
)
RECURSOS_PLANTILLA = (
    (
        "xl/media/image3.png",
        TwoCellAnchor(
            _from=AnchorMarker(col=0, colOff=243840, row=0, rowOff=114300),
            to=AnchorMarker(col=2, colOff=388620, row=0, rowOff=944880),
        ),
        ("logo_ops.png", "logo_izquierdo.png", "logo_soledad.png"),
    ),
    (
        "xl/media/image4.png",
        TwoCellAnchor(
            _from=AnchorMarker(col=12, colOff=350520, row=0, rowOff=76200),
            to=AnchorMarker(col=13, colOff=586740, row=0, rowOff=960120),
        ),
        ("logo_pjp4.png", "logo_derecho.png"),
    ),
    (
        "xl/media/image1.png",
        TwoCellAnchor(
            _from=AnchorMarker(col=2, colOff=883920, row=62, rowOff=76200),
            to=AnchorMarker(col=4, colOff=38100, row=66, rowOff=609600),
        ),
        ("firma_izquierda.png", "firma_luis.png"),
    ),
)
FIRMAS_ESTATICAS = ("firma_ops.png", "firma_derecha.png", "firma_gerardo.png")
_MARGEN_RANURA = 8
_FUENTE_PDF = "ReporteFoto"
_FUENTE_PDF_NEGRITA = "ReporteFoto-Bold"
CABECERA_PDF = (
    ("N° OT", "A4", 0.12),
    ("Ubicación", "C4", 0.13),
    ("Fecha", "D4", 0.12),
    ("Descripción de la Actividad", "E3", 0.37),
    ("% Avance", "K4", 0.10),
    ("Cuadrilla", "L4", 0.16),
)
FIRMANTES_PDF = (
    ("LUIS CARMELO CAMPOS BRITO", "COORDINADOR Y/O GERENTE DE PJP4"),
    ("GERARDO BERNABE RICARDEZ", "COORDINADOR DE OPS"),
)
COLOR_MARCO_RGB = (204, 0, 0)  # #CC0000
GROSOR_LINEA_MARCO = 2
SEPARACION_MARCO = 6
COLOR_BARRA_FIRMAS = colors.HexColor("#E5A200")
JPEG_CALIDAD_PDF = 85
JPEG_CALIDAD_EXCEL = 85
JPEG_DPI_PDF = 110
MAX_ANCHO_FOTO_PX = 1000
MAX_ALTO_FOTO_PX = 1000
MAX_LADO_FOTO_EXCEL_PX = 1000


def _ancho_columna_px(ws: Worksheet, columna: int) -> float:
    letra = get_column_letter(columna)
    ancho = ws.column_dimensions[letra].width
    if not ancho:
        ancho = ws.sheet_format.defaultColWidth or 8.43
    return ((256 * float(ancho) + int(128 / 7)) / 256) * 7


def _alto_fila_px(ws: Worksheet, fila: int) -> float:
    alto = ws.row_dimensions[fila].height
    if not alto:
        alto = 15
    return float(alto) * 96 / 72


def _origen_px(ws: Worksheet, columna: int, fila: int) -> tuple[float, float]:
    x = sum(_ancho_columna_px(ws, c) for c in range(1, columna))
    y = sum(_alto_fila_px(ws, r) for r in range(1, fila))
    return x, y


def _tamano_rango_px(ws: Worksheet, col_ini: int, fila_ini: int, col_fin: int, fila_fin: int) -> tuple[float, float]:
    ancho = sum(_ancho_columna_px(ws, c) for c in range(col_ini, col_fin + 1))
    alto = sum(_alto_fila_px(ws, r) for r in range(fila_ini, fila_fin + 1))
    return ancho, alto


def _ajustar_en_caja(ancho: int, alto: int, caja_ancho: float, caja_alto: float) -> tuple[int, int, int, int]:
    """Escala la imagen dentro de la caja y la centra. Conserva la proporción."""
    util_ancho = max(1, int(caja_ancho) - (2 * _MARGEN_RANURA))
    util_alto = max(1, int(caja_alto) - (2 * _MARGEN_RANURA))
    escala = min(util_ancho / max(ancho, 1), util_alto / max(alto, 1))
    nuevo_ancho = max(1, int(round(ancho * escala)))
    nuevo_alto = max(1, int(round(alto * escala)))
    desplazamiento_x = max(_MARGEN_RANURA, int((caja_ancho - nuevo_ancho) / 2))
    desplazamiento_y = max(_MARGEN_RANURA, int((caja_alto - nuevo_alto) / 2))
    return nuevo_ancho, nuevo_alto, desplazamiento_x, desplazamiento_y


def _a_rgb(imagen: PILImage.Image) -> PILImage.Image:
    if imagen.mode == "RGB":
        return imagen
    if imagen.mode == "RGBA":
        fondo = PILImage.new("RGB", imagen.size, "white")
        fondo.paste(imagen, mask=imagen.getchannel("A"))
        return fondo
    return imagen.convert("RGB")


def _redimensionar_maximo(imagen: PILImage.Image, max_lado: int) -> PILImage.Image:
    ancho, alto = imagen.size
    mayor = max(ancho, alto)
    if mayor <= max_lado:
        return imagen
    escala = max_lado / mayor
    return imagen.resize(
        (max(1, int(ancho * escala)), max(1, int(alto * escala))),
        PILImage.Resampling.LANCZOS,
    )


def _limitar_tamano_foto(imagen: PILImage.Image) -> PILImage.Image:
    """Si el ancho o el alto superan 1000 px, baja la imagen conservando proporción."""
    ancho, alto = imagen.size
    escala = min(1.0, MAX_ANCHO_FOTO_PX / max(ancho, 1), MAX_ALTO_FOTO_PX / max(alto, 1))
    if escala >= 1:
        return imagen
    return imagen.resize(
        (max(1, int(ancho * escala)), max(1, int(alto * escala))),
        PILImage.Resampling.LANCZOS,
    )


def _comprimir_foto_reporte(origen, calidad: int = JPEG_CALIDAD_PDF) -> io.BytesIO:
    """Abre la foto, la recorta a 1000 px y la deja como JPEG 85 en memoria."""
    if isinstance(origen, PILImage.Image):
        imagen = origen
        if imagen.mode not in {"RGB", "RGBA"}:
            imagen = imagen.convert("RGB")
        elif imagen.mode == "RGBA":
            imagen = _a_rgb(imagen)
        else:
            imagen = imagen.copy()
    else:
        with PILImage.open(origen) as crudo:
            imagen = ImageOps.exif_transpose(crudo)
            imagen = _a_rgb(imagen.copy())
    imagen = _limitar_tamano_foto(imagen)
    buffer = io.BytesIO()
    imagen.save(buffer, format="JPEG", quality=int(calidad), optimize=True)
    buffer.seek(0)
    buffer.name = f"evidencia_{id(buffer)}.jpg"
    return buffer


def _jpeg_bytes(imagen: PILImage.Image, calidad: int, max_lado: int | None = None) -> io.BytesIO:
    foto = _a_rgb(imagen)
    if max_lado:
        foto = _redimensionar_maximo(foto, max_lado)
    return _comprimir_foto_reporte(foto, calidad)


def _jpeg_para_pdf(imagen: PILImage.Image, ancho_pt: float, alto_pt: float) -> io.BytesIO:
    """Comprime la foto al tamaño de impresión y la entrega como JPEG en memoria."""
    objetivo_w = max(1, int(ancho_pt * JPEG_DPI_PDF / 72))
    objetivo_h = max(1, int(alto_pt * JPEG_DPI_PDF / 72))
    escala = min(objetivo_w / max(imagen.width, 1), objetivo_h / max(imagen.height, 1), 1.0)
    foto = _a_rgb(imagen)
    if escala < 0.98:
        foto = foto.resize(
            (max(1, int(imagen.width * escala)), max(1, int(imagen.height * escala))),
            PILImage.Resampling.LANCZOS,
        )
    return _comprimir_foto_reporte(foto, JPEG_CALIDAD_PDF)


def _preparar_imagen(ruta: str, max_ancho: int, max_alto: int) -> tuple[PILImage.Image, int, int]:
    with PILImage.open(_comprimir_foto_reporte(ruta, JPEG_CALIDAD_EXCEL)) as origen:
        imagen = origen.convert("RGB")
    ancho, alto = imagen.size
    nuevo_ancho, nuevo_alto, _, _ = _ajustar_en_caja(
        ancho, alto, max_ancho + (2 * _MARGEN_RANURA), max_alto + (2 * _MARGEN_RANURA)
    )
    if nuevo_ancho < ancho or nuevo_alto < alto:
        imagen = imagen.resize((nuevo_ancho, nuevo_alto), PILImage.Resampling.LANCZOS)
    return imagen, imagen.size[0], imagen.size[1]


def _es_metaarchivo(datos: bytes, nombre: str = "") -> bool:
    """WMF y EMF no tienen cargador en el Pillow de Linux."""
    extension = os.path.splitext(str(nombre or ""))[1].lower()
    if extension in {".wmf", ".emf"}:
        return True
    if len(datos) >= 4 and datos[:4] == b"\xd7\xcd\xc6\x9a":
        return True
    if len(datos) >= 44 and datos[40:44] == b" EMF":
        return True
    if len(datos) >= 4 and datos[:4] in {b"\x01\x00\x09\x00", b"\x02\x00\x09\x00"}:
        return True
    return False


def _bytes_raster(datos: bytes, nombre: str = "") -> bytes | None:
    """Devuelve PNG o JPEG. Si Pillow no puede leer el WMF/EMF, regresa None."""
    if not datos:
        return None
    try:
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            with PILImage.open(io.BytesIO(datos)) as imagen:
                imagen.load()
                formato = str(imagen.format or "").upper()
                if formato in {"PNG", "JPEG", "JPG", "GIF"} and not _es_metaarchivo(datos, nombre):
                    return datos
                buffer = io.BytesIO()
                imagen.convert("RGB").save(buffer, format="PNG")
                return buffer.getvalue()
    except Exception:
        return None


def _archivo_estatico(*nombres: str) -> bytes | None:
    carpetas = (
        os.path.dirname(PLANTILLA_XLSX) or "database",
        os.path.join("database", "logos"),
        "static",
        os.path.join("static", "logos"),
    )
    for carpeta in carpetas:
        for nombre in nombres:
            ruta = os.path.join(carpeta, nombre)
            if not os.path.isfile(ruta):
                continue
            try:
                with open(ruta, "rb") as archivo:
                    datos = archivo.read()
            except OSError:
                continue
            raster = _bytes_raster(datos, ruta)
            if raster:
                return raster
    return None


def _png_minimo() -> bytes:
    buffer = io.BytesIO()
    PILImage.new("RGBA", (1, 1), (0, 0, 0, 0)).save(buffer, format="PNG")
    return buffer.getvalue()


def _plantilla_sin_vectores(ruta: str) -> io.BytesIO:
    """Sustituye WMF/EMF por un PNG vacío para que openpyxl pueda abrir el libro."""
    reemplazo = _png_minimo()
    salida = io.BytesIO()
    with zipfile.ZipFile(ruta) as origen, zipfile.ZipFile(salida, "w") as destino:
        for info in origen.infolist():
            datos = origen.read(info.filename)
            if _es_metaarchivo(datos, info.filename):
                datos = reemplazo
            destino.writestr(info.filename, datos)
    salida.seek(0)
    return salida


def _buffer_png(datos: bytes, nombre: str) -> io.BytesIO:
    buffer = io.BytesIO(datos)
    buffer.name = nombre
    return buffer


def _firma_derecha_png(ruta_plantilla: str) -> io.BytesIO | None:
    estatica = _archivo_estatico(*FIRMAS_ESTATICAS)
    if estatica:
        return _buffer_png(estatica, "firma_derecha.png")
    datos = b""
    try:
        with zipfile.ZipFile(ruta_plantilla) as paquete:
            for nombre in paquete.namelist():
                if os.path.splitext(nombre)[1].lower() in {".emf", ".wmf"}:
                    datos = paquete.read(nombre)
                    break
    except (FileNotFoundError, zipfile.BadZipFile, KeyError):
        datos = b""
    raster = _bytes_raster(datos, "firma.emf") if datos else None
    if not raster:
        return None
    return _buffer_png(raster, "firma_derecha.png")


def _enmarcar_foto(
    imagen: PILImage.Image,
    linea: int = GROSOR_LINEA_MARCO,
    separacion: int = SEPARACION_MARCO,
) -> PILImage.Image:
    """Línea roja delgada con un margen blanco entre la foto y el marco."""
    if imagen.mode != "RGB":
        imagen = imagen.convert("RGB")
    linea = max(1, int(linea))
    separacion = max(1, int(separacion))
    margen = linea + separacion
    marco = PILImage.new("RGB", (imagen.width + 2 * margen, imagen.height + 2 * margen), (255, 255, 255))
    marco.paste(imagen, (margen, margen))
    ImageDraw.Draw(marco).rectangle(
        (0, 0, marco.width - 1, marco.height - 1),
        outline=COLOR_MARCO_RGB,
        width=linea,
    )
    return marco


def _anclar_foto(ws: Worksheet, imagen: PILImage.Image, ranura: tuple[int, int, int, int]) -> dict:
    col_ini, fila_ini, col_fin, fila_fin = ranura
    caja_ancho, caja_alto = _tamano_rango_px(ws, col_ini, fila_ini, col_fin, fila_fin)
    ancho, alto, dx, dy = _ajustar_en_caja(imagen.size[0], imagen.size[1], caja_ancho, caja_alto)
    margen = GROSOR_LINEA_MARCO + SEPARACION_MARCO
    if ancho > 2 * margen + 8 and alto > 2 * margen + 8:
        escala = min(
            (ancho - 2 * margen) / max(imagen.size[0], 1),
            (alto - 2 * margen) / max(imagen.size[1], 1),
        )
        foto_ancho = max(1, int(round(imagen.size[0] * escala)))
        foto_alto = max(1, int(round(imagen.size[1] * escala)))
        foto = imagen
        if (foto_ancho, foto_alto) != imagen.size:
            foto = imagen.resize((foto_ancho, foto_alto), PILImage.Resampling.LANCZOS)
        imagen = _enmarcar_foto(foto)
        ancho, alto = imagen.size
        dx = max(_MARGEN_RANURA, int((caja_ancho - ancho) / 2))
        dy = max(_MARGEN_RANURA, int((caja_alto - alto) / 2))
    elif (ancho, alto) != imagen.size:
        imagen = imagen.resize((ancho, alto), PILImage.Resampling.LANCZOS)
    imagen = _redimensionar_maximo(imagen, MAX_LADO_FOTO_EXCEL_PX)
    ancho, alto = imagen.size
    buffer = _jpeg_bytes(imagen, JPEG_CALIDAD_EXCEL, MAX_LADO_FOTO_EXCEL_PX)
    dibujo = XLImage(buffer)
    dibujo.width = ancho
    dibujo.height = alto
    dibujo.anchor = OneCellAnchor(
        _from=AnchorMarker(
            col=col_ini - 1,
            colOff=pixels_to_EMU(dx),
            row=fila_ini - 1,
            rowOff=pixels_to_EMU(dy),
        ),
        ext=XDRPositiveSize2D(pixels_to_EMU(ancho), pixels_to_EMU(alto)),
    )
    ws.add_image(dibujo)
    origen_x, origen_y = _origen_px(ws, col_ini, fila_ini)
    return {
        "x": origen_x + dx,
        "y": origen_y + dy,
        "ancho": ancho,
        "alto": alto,
        "ranura": (
            origen_x,
            origen_y,
            caja_ancho,
            caja_alto,
        ),
        "proporcion": ancho / alto,
        "imagen": imagen,
    }


def _quitar_metaarchivos(hoja: Worksheet) -> None:
    conservadas = []
    for imagen in list(getattr(hoja, "_images", []) or []):
        formato = str(getattr(imagen, "format", "") or "").lower()
        ruta = str(getattr(imagen, "path", "") or "")
        if formato in {"wmf", "emf"} or _es_metaarchivo(b"", ruta):
            continue
        conservadas.append(imagen)
    hoja._images = conservadas


def _recursos_oficiales(ruta_plantilla: str) -> list[tuple[bytes, TwoCellAnchor]]:
    recursos = []
    try:
        paquete = zipfile.ZipFile(ruta_plantilla)
    except (FileNotFoundError, zipfile.BadZipFile):
        paquete = None
    for nombre, ancla, respaldos in RECURSOS_PLANTILLA:
        raster = None
        if paquete is not None:
            try:
                raster = _bytes_raster(paquete.read(nombre), nombre)
            except KeyError:
                raster = None
        if raster is None:
            raster = _archivo_estatico(*respaldos)
        if raster:
            recursos.append((raster, ancla))
    if paquete is not None:
        paquete.close()
    return recursos


def _clonar_ancla(ancla: TwoCellAnchor) -> TwoCellAnchor:
    return TwoCellAnchor(
        _from=AnchorMarker(
            col=ancla._from.col,
            colOff=ancla._from.colOff,
            row=ancla._from.row,
            rowOff=ancla._from.rowOff,
        ),
        to=AnchorMarker(
            col=ancla.to.col,
            colOff=ancla.to.colOff,
            row=ancla.to.row,
            rowOff=ancla.to.rowOff,
        ),
    )


def _colocar_recurso(ws: Worksheet, datos: bytes, ancla: TwoCellAnchor, retenidas: list, nombre: str) -> None:
    raster = _bytes_raster(datos, nombre)
    if not raster:
        return
    try:
        buffer = _buffer_png(raster, nombre if nombre.lower().endswith((".png", ".jpg", ".jpeg")) else "recurso.png")
        imagen = XLImage(buffer)
        imagen.anchor = _clonar_ancla(ancla)
        ws.add_image(imagen)
        retenidas.append(buffer)
    except Exception:
        return


def _agregar_firma(ws: Worksheet, firma: io.BytesIO | None, retenidas: list) -> None:
    if firma is None:
        return
    try:
        firma.seek(0)
        copia = io.BytesIO(firma.read())
        copia.name = "firma_derecha.png"
        imagen = XLImage(copia)
        imagen.anchor = _clonar_ancla(FIRMA_DERECHA)
        ws.add_image(imagen)
        retenidas.append(copia)
    except Exception:
        return


def _campos_ubicacion_actividad(actividad: dict) -> tuple:
    return (
        actividad.get("ubicacion"),
        actividad.get("pozo"),
        actividad.get("sector"),
        actividad.get("equipo"),
        actividad.get("tag"),
    )


def _escribir_encabezado(ws: Worksheet, actividad: dict, continuacion: str) -> None:
    campos = _campos_ubicacion_actividad(actividad)
    ubicacion = _ubicacion_tabla_pdf(*campos)
    descripcion = _descripcion_tabla_pdf(actividad.get("descripcion") or "", ubicacion, *campos)
    if continuacion:
        descripcion = f"{descripcion} {continuacion}".strip()
    ws["A4"] = str(actividad.get("ot") or "")
    ws["C4"] = ubicacion or None
    fecha = actividad.get("fecha")
    if isinstance(fecha, datetime):
        fecha = fecha.date()
    ws["D4"] = fecha if isinstance(fecha, date) else str(fecha or "")
    if isinstance(fecha, date):
        ws["D4"].number_format = "DD/MM/YYYY"
    ws["E3"] = descripcion
    ws["K4"] = _texto_avance_pdf(actividad.get("avance"))
    ws["L4"] = str(actividad.get("cuadrilla") or "")
    for coordenada in ("C4", "E3", "L4"):
        ws[coordenada].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _nombre_hoja(indice: int, actividad: dict, parte: int, usados: set[str]) -> str:
    fecha = actividad.get("fecha")
    dia = fecha.strftime("%d%m") if isinstance(fecha, date) else ""
    cuadrilla = re.sub(r"[:\\/?*\[\]]+", " ", str(actividad.get("cuadrilla") or "Cuadrilla"))
    cuadrilla = re.sub(r"\s+", " ", cuadrilla).strip()
    sufijo = f" {parte}" if parte else ""
    base = f"{indice:02d} {cuadrilla} {dia}{sufijo}".strip()
    base = base[:31] or f"Hoja {indice}"
    nombre = base
    extra = 2
    while nombre.casefold() in usados:
        recorte = base[: 31 - len(str(extra)) - 1]
        nombre = f"{recorte} {extra}"
        extra += 1
    usados.add(nombre.casefold())
    return nombre


def _preparar_libro(ruta_plantilla: str) -> tuple[Workbook, Worksheet, io.BytesIO | None]:
    if not os.path.isfile(ruta_plantilla):
        raise FileNotFoundError(
            f"No está la plantilla del reporte fotográfico en {ruta_plantilla}."
        )
    with warnings.catch_warnings():
        warnings.filterwarnings(
            "ignore",
            message="wmf image format is not supported",
            category=UserWarning,
        )
        warnings.filterwarnings(
            "ignore",
            message="will be removed because it cannot be read",
            category=UserWarning,
        )
        try:
            libro = load_workbook(ruta_plantilla)
        except Exception:
            libro = load_workbook(_plantilla_sin_vectores(ruta_plantilla))
            for hoja in libro.worksheets:
                hoja._images = [
                    imagen
                    for imagen in list(getattr(hoja, "_images", []) or [])
                    if not (getattr(imagen, "width", None) == 1 and getattr(imagen, "height", None) == 1)
                ]
    libro._external_links = []
    for hoja in libro.worksheets:
        _quitar_metaarchivos(hoja)
    plantilla = libro.active
    plantilla["K4"] = plantilla["K4"].value if not isinstance(plantilla["K4"].value, str) or not str(plantilla["K4"].value).startswith("=") else ""
    if isinstance(plantilla["Q21"].value, str) and str(plantilla["Q21"].value).startswith("="):
        plantilla["Q21"] = None
    firma = _firma_derecha_png(ruta_plantilla)
    return libro, plantilla, firma


def _dibujar_rubrica(pdf, lector, x: float, y: float, ancho: float, alto: float) -> None:
    """Apoya la rúbrica sobre la línea, centrada y sin deformarla."""
    try:
        imagen_ancho, imagen_alto = lector.getSize()
    except Exception:
        return
    if not imagen_ancho or not imagen_alto:
        return
    escala = min(ancho / imagen_ancho, alto / imagen_alto)
    dibujo_ancho = imagen_ancho * escala
    dibujo_alto = imagen_alto * escala
    pdf.drawImage(
        lector,
        x + (ancho - dibujo_ancho) / 2,
        y,
        width=dibujo_ancho,
        height=dibujo_alto,
        preserveAspectRatio=True,
        mask="auto",
    )


def _dibujar_bloque_firmas(
    pdf: canvas.Canvas,
    ws: Worksheet,
    tabla_x: float,
    tabla_ancho: float,
    firmas: list[tuple[bytes, int]],
    fuente: str,
    fuente_negrita: str,
    y_fotos: float,
) -> None:
    """Barra amarilla separada y recuadro bajo. El blanco queda al pie de la hoja."""
    separacion_fotos = 20.0
    alto_barra = 7.0
    separacion_barra = 12.0
    alto_caja = 58.0
    separacion_nombre = 16.0
    separacion_cargo = 13.0
    margen_inferior = 56.0
    y_barra = y_fotos - separacion_fotos - alto_barra
    y_tope = y_barra - separacion_barra
    y_caja = y_tope - alto_caja
    y_nombre = y_caja - separacion_nombre
    y_cargo = y_nombre - separacion_cargo
    if y_cargo < margen_inferior:
        libre = y_fotos - (y_barra + alto_barra)
        ajuste = min(margen_inferior - y_cargo, max(0.0, libre - 12.0))
        y_barra += ajuste
        y_tope += ajuste
        y_caja += ajuste
        y_nombre += ajuste
        y_cargo += ajuste
    y_leyenda = y_tope - 12.0
    y_imagen = y_caja + 3.0
    alto_imagen = max(28.0, y_leyenda - 5.0 - y_imagen)
    por_columna: dict[int, bytes] = {}
    for datos, columna in firmas:
        por_columna[0 if columna < 7 else 1] = datos
    pdf.setFillColor(COLOR_BARRA_FIRMAS)
    pdf.rect(tabla_x, y_barra, tabla_ancho, alto_barra, stroke=0, fill=1)
    pdf.setStrokeColorRGB(0, 0, 0)
    pdf.setLineWidth(1)
    pdf.rect(tabla_x, y_caja, tabla_ancho, alto_caja, stroke=1, fill=0)
    mitad = tabla_x + tabla_ancho / 2
    pdf.setLineWidth(0.7)
    pdf.line(mitad, y_caja, mitad, y_tope)
    for indice, (nombre, cargo) in enumerate(FIRMANTES_PDF):
        x_columna = tabla_x + indice * tabla_ancho / 2
        ancho_columna = tabla_ancho / 2
        centro = x_columna + ancho_columna / 2
        margen = 8.0
        datos = por_columna.get(indice)
        if datos:
            try:
                _dibujar_rubrica(
                    pdf,
                    ImageReader(io.BytesIO(datos)),
                    x_columna + margen,
                    y_imagen,
                    ancho_columna - 2 * margen,
                    alto_imagen,
                )
            except Exception:
                pass
        etiqueta = _texto_celda(ws["A62" if indice == 0 else "H62"].value) or "APROBADO:"
        pdf.setFillColorRGB(0.1, 0.14, 0.12)
        pdf.setFont(fuente, 7)
        pdf.drawCentredString(centro, y_leyenda, etiqueta)
        pdf.setFont(fuente_negrita, 8)
        pdf.drawCentredString(centro, y_nombre, nombre)
        pdf.setFont(fuente, 7)
        pdf.drawCentredString(centro, y_cargo, cargo)


def _clave_foto(ruta: str) -> str:
    texto = str(ruta or "").strip()
    if not texto:
        return ""
    if os.path.isfile(texto):
        return os.path.normcase(os.path.abspath(texto))
    return os.path.normcase(texto)


def _bloques_fotos(rutas: list[str]) -> list[list[str]]:
    """Cada página recibe solo fotos que todavía no se usaron en la actividad."""
    vistas: set[str] = set()
    unicas: list[str] = []
    for ruta in rutas:
        clave = _clave_foto(ruta)
        if not clave or clave in vistas:
            continue
        vistas.add(clave)
        unicas.append(str(ruta).strip())
    if not unicas:
        return [[]]
    return [list(unicas[i : i + FOTOS_POR_HOJA]) for i in range(0, len(unicas), FOTOS_POR_HOJA)]


def _limpiar_fotos_cuadricula(hoja: Worksheet) -> None:
    """Quita fotos de evidencia heredadas antes de armar esta página."""
    conservadas = []
    for imagen in list(getattr(hoja, "_images", []) or []):
        ancla = getattr(imagen, "anchor", None)
        origen = getattr(ancla, "_from", None)
        fila = getattr(origen, "row", None)
        if isinstance(fila, int) and 8 <= fila <= 52:
            continue
        conservadas.append(imagen)
    hoja._images = conservadas


def _fotos_de_bloque(hoja: Worksheet, bloque: list[str]) -> tuple[list[tuple[PILImage.Image, dict]], int]:
    """Carga solo el bloque de esta página, en una lista nueva."""
    fotos_hoja: list[tuple[PILImage.Image, dict]] = []
    omitidas = 0
    for posicion, ruta in enumerate(list(bloque)[:FOTOS_POR_HOJA]):
        if not os.path.isfile(ruta):
            omitidas += 1
            continue
        ranura = RANURAS[posicion]
        caja_ancho, caja_alto = _tamano_rango_px(hoja, *ranura)
        try:
            imagen, _, _ = _preparar_imagen(
                ruta,
                max(1, int(caja_ancho) - 2 * _MARGEN_RANURA),
                max(1, int(caja_alto) - 2 * _MARGEN_RANURA),
            )
        except Exception:
            omitidas += 1
            continue
        caja = _anclar_foto(hoja, imagen, ranura)
        caja["dx"] = caja["x"] - caja["ranura"][0]
        caja["dy"] = caja["y"] - caja["ranura"][1]
        caja["ranura_idx"] = ranura
        propia = caja["imagen"].copy()
        caja["imagen"] = propia
        caja["origen"] = (propia.size[0], propia.size[1])
        fotos_hoja.append((propia, caja))
    return fotos_hoja, omitidas


def _hoja_a_pdf(
    pdf: canvas.Canvas,
    ws: Worksheet,
    fotos: list[tuple[PILImage.Image, dict]],
    recursos: list[tuple[bytes, TwoCellAnchor]],
) -> None:
    fotos = list(fotos or [])
    ancho_hoja, alto_hoja = letter
    margen = 8
    ancho_contenido = sum(_ancho_columna_px(ws, c) for c in range(1, 15)) * 72 / 96
    alto_contenido = sum((ws.row_dimensions[r].height or 15) for r in range(1, 70))
    escala = min(
        (ancho_hoja - 2 * margen) / max(ancho_contenido, 1),
        (alto_hoja - 2 * margen) / max(alto_contenido, 1),
    )
    fuente, fuente_negrita = _fuentes_pdf()

    def x_pt(columna: int, extra: float = 0) -> float:
        return margen + (sum(_ancho_columna_px(ws, c) for c in range(1, columna)) + extra) * 72 / 96 * escala

    def y_pt(fila: int, extra: float = 0) -> float:
        acumulado = sum((ws.row_dimensions[r].height or 15) for r in range(1, fila)) + extra * 72 / 96
        return alto_hoja - margen - acumulado * escala

    def alto_pt(fila_ini: int, fila_fin: int) -> float:
        return sum((ws.row_dimensions[r].height or 15) for r in range(fila_ini, fila_fin + 1)) * escala

    def ancho_pt(col_ini: int, col_fin: int) -> float:
        return sum(_ancho_columna_px(ws, c) for c in range(col_ini, col_fin + 1)) * 72 / 96 * escala

    pdf.setFillColorRGB(1, 1, 1)
    pdf.rect(0, 0, ancho_hoja, alto_hoja, stroke=0, fill=1)

    firmas: list[tuple[bytes, int]] = []
    for datos, ancla in recursos:
        if ancla._from.row >= 50:
            firmas.append((datos, ancla._from.col))
            continue
        try:
            lector = ImageReader(io.BytesIO(datos))
        except Exception:
            continue
        x0 = x_pt(ancla._from.col + 1, ancla._from.colOff / 9525)
        y_sup = y_pt(ancla._from.row + 1, ancla._from.rowOff / 9525)
        x1 = x_pt(ancla.to.col + 1, ancla.to.colOff / 9525)
        y_inf = y_pt(ancla.to.row + 1, ancla.to.rowOff / 9525)
        _dibujar_proporcion(pdf, lector, x0, y_inf, max(1, x1 - x0), max(1, y_sup - y_inf))

    pdf.setFillColorRGB(0.08, 0.12, 0.1)
    pdf.setFont(fuente_negrita, 11)
    pdf.drawCentredString(x_pt(4) + ancho_pt(4, 12) / 2, y_pt(1) - 28, "REPORTE DE EVIDENCIA FOTOGRAFICA")
    tabla_x = x_pt(1)
    tabla_ancho = ancho_pt(1, 14)
    tabla_arriba = y_pt(3)
    tabla_abajo = y_pt(8)
    alto_etiquetas = min(16.0, (tabla_arriba - tabla_abajo) * 0.3)
    alto_valores = tabla_arriba - tabla_abajo - alto_etiquetas
    anchos = [tabla_ancho * proporcion for _etiqueta, _celda, proporcion in CABECERA_PDF]
    padding = 2
    fila_etiquetas = [
        _parrafo_ajustado(etiqueta, fuente_negrita, 7.5, ancho - 2 * padding, alto_etiquetas - 2)
        for (etiqueta, _celda, _p), ancho in zip(CABECERA_PDF, anchos)
    ]
    fila_valores = [
        _parrafo_ajustado(
            _valor_tabla_pdf(celda, ws[celda].value, ws["C4"].value),
            fuente_negrita if celda == "E3" else fuente,
            7 if celda == "E3" else 8,
            ancho - 2 * padding,
            alto_valores - 2 * padding,
        )
        for (_etiqueta, celda, _p), ancho in zip(CABECERA_PDF, anchos)
    ]
    tabla = Table([fila_etiquetas, fila_valores], colWidths=anchos, rowHeights=[alto_etiquetas, alto_valores])
    tabla.setStyle(
        TableStyle(
            [
                ("BOX", (0, 0), (-1, -1), 1.2, colors.black),
                ("INNERGRID", (0, 0), (-1, -1), 0.8, colors.black),
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E6E6E6")),
                ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("LEFTPADDING", (0, 0), (-1, -1), padding),
                ("RIGHTPADDING", (0, 0), (-1, -1), padding),
                ("TOPPADDING", (0, 0), (-1, -1), 1),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 1),
            ]
        )
    )
    tabla.wrapOn(pdf, tabla_ancho, tabla_arriba - tabla_abajo)
    tabla.drawOn(pdf, tabla_x, tabla_abajo)

    margen_pt = _MARGEN_RANURA * 72 / 96 * escala
    cajas_ranura = [ancho_pt(c0, c1) for c0, _f0, c1, _f1 in RANURAS]
    altos_ranura = [alto_pt(f0, f1) for _c0, f0, _c1, f1 in RANURAS]
    ancho_fijo = max(1.0, min(cajas_ranura) - 2 * margen_pt)
    alto_fijo = max(1.0, min(altos_ranura) - 2 * margen_pt)
    for imagen, caja in fotos:
        col_ini, fila_ini, col_fin, fila_fin = _ranura_de_caja(ws, caja)
        slot_ancho = ancho_pt(col_ini, col_fin)
        slot_alto = alto_pt(fila_ini, fila_fin)
        _dibujar_foto_uniforme(
            pdf,
            ImageReader(_jpeg_para_pdf(imagen, ancho_fijo, alto_fijo)),
            x_pt(col_ini) + (slot_ancho - ancho_fijo) / 2,
            y_pt(fila_fin + 1) + (slot_alto - alto_fijo) / 2,
            ancho_fijo,
            alto_fijo,
        )

    _dibujar_bloque_firmas(
        pdf,
        ws,
        tabla_x,
        tabla_ancho,
        firmas,
        fuente,
        fuente_negrita,
        y_pt(54),
    )
    pdf.showPage()


def _ranura_de_caja(ws: Worksheet, caja: dict) -> tuple[int, int, int, int]:
    return caja["ranura_idx"]


def _caja_proporcion(lector, x, y, ancho, alto) -> tuple[float, float, float, float] | None:
    try:
        imagen_ancho, imagen_alto = lector.getSize()
    except Exception:
        return None
    if not imagen_ancho or not imagen_alto:
        return None
    escala = min(ancho / imagen_ancho, alto / imagen_alto)
    dibujo_ancho = imagen_ancho * escala
    dibujo_alto = imagen_alto * escala
    return (
        x + (ancho - dibujo_ancho) / 2,
        y + (alto - dibujo_alto) / 2,
        dibujo_ancho,
        dibujo_alto,
    )


def _dibujar_foto_uniforme(pdf, lector, x: float, y: float, ancho: float, alto: float) -> None:
    """Caja fija de la cuadrícula: conserva la proporción y centra la foto."""
    pdf.drawImage(
        lector,
        x,
        y,
        width=ancho,
        height=alto,
        preserveAspectRatio=True,
        anchor="c",
        mask="auto",
    )


def _dibujar_proporcion(pdf, lector, x, y, ancho, alto) -> tuple[float, float, float, float] | None:
    caja = _caja_proporcion(lector, x, y, ancho, alto)
    if caja is None:
        pdf.drawImage(lector, x, y, width=ancho, height=alto, preserveAspectRatio=True, mask="auto", anchor="c")
        return None
    x_dibujo, y_dibujo, dibujo_ancho, dibujo_alto = caja
    pdf.drawImage(
        lector,
        x_dibujo,
        y_dibujo,
        width=dibujo_ancho,
        height=dibujo_alto,
        preserveAspectRatio=True,
        mask="auto",
    )
    return caja


def _texto_celda(valor) -> str:
    if valor is None:
        return ""
    if isinstance(valor, datetime):
        return valor.strftime("%d/%m/%Y")
    if isinstance(valor, date):
        return valor.strftime("%d/%m/%Y")
    return str(valor).strip()


def _sin_acentos(texto: str) -> str:
    normal = unicodedata.normalize("NFD", str(texto or ""))
    return "".join(c for c in normal if unicodedata.category(c) != "Mn")


_TAGS_CUADRILLA_PDF = {
    "electricos": "ELECT",
    "electrico": "ELECT",
    "instrumentos": "INSTR",
    "instrumentista": "INSTR",
    "mecanico 1": "MECA-1",
    "mecanico 2": "MECA-2",
    "mecanico 3": "MECA-3",
    "mecanico 4": "MECA-4",
    "montaje": "MONTA",
    "soldadura": "SOLD",
    "retro": "RETRO",
    "retroexcavadora": "RETRO",
    "retro excavadora": "RETRO",
}


def _es_sector_zona(texto: str) -> bool:
    """El primer tramo de la ubicación de campo es el sector, no el pozo."""
    limpio = str(texto or "").strip()
    if not limpio:
        return False
    if "(" in limpio and ")" in limpio:
        return True
    clave = re.sub(r"\s+", " ", _sin_acentos(limpio).casefold()).strip()
    return clave in {"base pjp4"} or clave.startswith(("almacen ", "bateria ", "estacion ", "sector "))


_RE_NINGUNO_TEXTO = re.compile(r"\bninguno\b", re.I)


def _sin_ninguno_pdf(valor) -> str:
    texto = sitios_operativos.quitar_ninguno(valor)
    texto = _RE_NINGUNO_TEXTO.sub(" ", texto)
    texto = re.sub(r"\s+", " ", texto).strip(" ·|-/,;")
    if not texto or sitios_operativos.es_vacio_sitio(texto) or texto.upper() == "NINGUNO":
        return ""
    return texto


def _pozo_tabla_pdf(valor) -> str:
    """Nombre de la instalación o clave del pozo. Nunca 'Ninguno'."""
    return _ubicacion_tabla_pdf(valor)


_RE_TAG_POZO_PDF = re.compile(r"\b((?:PSN|PGA|PPB|PS|PA|PB|PC)-?\d+[A-Z]?)\b", re.I)
_RE_INFRA_PDF = re.compile(r"\b((?:GDO|CBN|OLD|OLG)-[A-Z0-9]+-IQ-SIN-\d+)\b", re.I)


def _compactar_tag_pozo_pdf(tag: str) -> str:
    texto = str(tag or "").strip().upper()
    return re.sub(r"(?<=[A-Z])-(?=\d)", "", texto)


def _ubicacion_tabla_pdf(*valores) -> str:
    """Columna Ubicación: solo el tag corto (PSN77, BSNI, ECSN). Sin macrozona ni MCI."""
    clave = sitios_operativos.clave_sitio_campos(*valores)
    if clave and not sitios_operativos.es_vacio_sitio(clave):
        return clave
    unidos = " ".join(_sin_ninguno_pdf(v) for v in valores if v)
    pozo = _RE_TAG_POZO_PDF.search(unidos)
    if pozo:
        return _compactar_tag_pozo_pdf(pozo.group(1))
    infra = _RE_INFRA_PDF.search(unidos.replace(" ", ""))
    if not infra:
        infra = _RE_INFRA_PDF.search(unidos)
    if infra:
        return infra.group(1).upper()
    for valor in valores:
        limpio = _sin_ninguno_pdf(valor)
        if not limpio:
            continue
        clave = sitios_operativos.clave_sitio(limpio)
        if clave:
            return clave
        if sitios_operativos.es_equipo_explicito(limpio):
            continue
        primero = re.split(r"\s*[-·|/]+\s*", limpio)[0].strip()
        pozo = _RE_TAG_POZO_PDF.match(primero)
        if pozo:
            return _compactar_tag_pozo_pdf(pozo.group(1))
        if " " in limpio or len(limpio) > 24:
            continue
        return limpio
    return ""


def _tag_cuadrilla_pdf(valor) -> str:
    texto = _texto_celda(valor)
    clave = re.sub(r"\s+", " ", _sin_acentos(texto).casefold()).strip()
    if clave in _TAGS_CUADRILLA_PDF:
        return _TAGS_CUADRILLA_PDF[clave]
    compacta = clave.replace(" ", "")
    if compacta in _TAGS_CUADRILLA_PDF:
        return _TAGS_CUADRILLA_PDF[compacta]
    return texto.upper()


def _descripcion_tabla_pdf(descripcion, *ubicaciones) -> str:
    """Agrega la clave del sitio al final. Nunca escribe EN NINGUNO."""
    texto = _sin_ninguno_pdf(descripcion)
    clave = sitios_operativos.clave_sitio_campos(*ubicaciones)
    if not clave:
        tag = sitios_operativos.clave_o_nombre_sitio(" · ".join(str(u or "") for u in ubicaciones if u))
        tag = _sin_ninguno_pdf(tag)
        if tag and " " not in tag and tag.upper() != "NINGUNO":
            clave = tag
    if not clave or sitios_operativos.es_vacio_sitio(clave) or _RE_NINGUNO_TEXTO.search(clave):
        return texto
    sufijo = f"EN {clave}"
    if _RE_NINGUNO_TEXTO.search(sufijo):
        return texto
    if texto.upper().endswith(sufijo.upper()):
        return texto
    if not texto:
        return sufijo
    return f"{texto} {sufijo}"


def _texto_avance_pdf(valor) -> str:
    texto = _texto_celda(valor).replace("%", "").strip()
    if not texto:
        return ""
    try:
        numero = float(texto.replace(",", "."))
    except ValueError:
        return _texto_celda(valor)
    if 0 < numero < 1:
        numero *= 100
    return f"{int(round(numero))}%"


def _valor_tabla_pdf(celda: str, valor, ubicacion=None) -> str:
    if celda == "K4":
        return _texto_avance_pdf(valor)
    if celda == "E3":
        return _descripcion_tabla_pdf(valor, ubicacion)
    if celda == "C4":
        return _ubicacion_tabla_pdf(valor)
    if celda == "L4":
        return _tag_cuadrilla_pdf(valor)
    return _sin_ninguno_pdf(_texto_celda(valor)) or _texto_celda(valor)


def _parrafo_ajustado(texto: str, fuente: str, tamano: float, ancho: float, alto: float) -> Paragraph:
    contenido = escape(texto or "")
    while True:
        estilo = ParagraphStyle(
            "celda",
            fontName=fuente,
            fontSize=tamano,
            leading=tamano + 1.5,
            alignment=TA_CENTER,
            textColor=colors.HexColor("#141F1A"),
        )
        parrafo = Paragraph(contenido, estilo)
        _ancho, alto_texto = parrafo.wrap(ancho, alto)
        if alto_texto <= alto or tamano <= 5:
            return parrafo
        tamano -= 0.5


def _fuentes_pdf() -> tuple[str, str]:
    if _FUENTE_PDF in pdfmetrics.getRegisteredFontNames():
        return _FUENTE_PDF, _FUENTE_PDF_NEGRITA
    windir = os.environ.get("WINDIR", r"C:\Windows")
    pares = (
        (
            os.path.join(windir, "Fonts", "arial.ttf"),
            os.path.join(windir, "Fonts", "arialbd.ttf"),
        ),
        (
            "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
            "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf",
        ),
        (
            "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf",
            "/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf",
        ),
    )
    for regular, negrita in pares:
        if os.path.isfile(regular) and os.path.isfile(negrita):
            pdfmetrics.registerFont(TTFont(_FUENTE_PDF, regular))
            pdfmetrics.registerFont(TTFont(_FUENTE_PDF_NEGRITA, negrita))
            return _FUENTE_PDF, _FUENTE_PDF_NEGRITA
    return "Helvetica", "Helvetica-Bold"


def _ordenar(actividades: list[dict]) -> list[dict]:
    def clave(item: dict):
        fecha = item.get("fecha") or date.min
        if isinstance(fecha, datetime):
            fecha = fecha.date()
        orden = item.get("orden")
        return (fecha, 99 if orden is None else int(orden), int(item.get("secuencia") or 0))

    return sorted(actividades, key=clave)


def generar_reporte_fotografico(
    actividades: list[dict],
    ruta_plantilla: str = PLANTILLA_XLSX,
) -> tuple[bytes, bytes, dict]:
    """Devuelve el Excel, el PDF y un resumen de hojas, fotos y omisiones."""
    import importlib

    importlib.reload(sitios_operativos)
    if not actividades:
        raise ValueError("No hay actividades en el periodo seleccionado.")
    libro, plantilla, firma = _preparar_libro(ruta_plantilla)
    recursos = _recursos_oficiales(ruta_plantilla)
    retenidas: list = []
    if firma is not None:
        retenidas.append(firma)
    usados: set[str] = set()
    paginas_pdf: list[tuple[Worksheet, list]] = []
    fotos_insertadas = 0
    fotos_omitidas = 0
    hojas = 0
    for indice, actividad in enumerate(_ordenar(actividades), start=1):
        rutas = [str(ruta) for ruta in (actividad.get("fotos") or []) if str(ruta).strip()]
        bloques = _bloques_fotos(rutas)
        for parte, bloque in enumerate(bloques):
            hoja = libro.copy_worksheet(plantilla)
            _quitar_metaarchivos(hoja)
            hoja.title = _nombre_hoja(indice, actividad, parte + 1 if len(bloques) > 1 else 0, usados)
            _escribir_encabezado(
                hoja,
                actividad,
                "(continuación)" if parte else "",
            )
            hoja.page_setup.orientation = "portrait"
            hoja.page_setup.fitToWidth = 1
            hoja.page_setup.fitToHeight = 1
            hoja.page_setup.paperSize = hoja.PAPERSIZE_LETTER
            hoja.sheet_properties.pageSetUpPr.fitToPage = True
            hoja.print_area = "A1:N69"
            for indice_recurso, (datos, ancla) in enumerate(recursos):
                _colocar_recurso(hoja, datos, ancla, retenidas, f"recurso_{indice_recurso}.png")
            _agregar_firma(hoja, firma, retenidas)
            _limpiar_fotos_cuadricula(hoja)
            fotos_hoja, omitidas_bloque = _fotos_de_bloque(hoja, list(bloque))
            fotos_omitidas += omitidas_bloque
            fotos_insertadas += len(fotos_hoja)
            paginas_pdf.append((hoja, list(fotos_hoja)))
            fotos_hoja.clear()
            hojas += 1
    libro.remove(plantilla)
    recursos_pdf = list(recursos)
    if firma is not None:
        firma.seek(0)
        recursos_pdf.append((firma.read(), FIRMA_DERECHA))
    pdf_buffer = io.BytesIO()
    pdf = canvas.Canvas(pdf_buffer, pagesize=letter)
    pdf.setPageCompression(1)
    for hoja, fotos_hoja in paginas_pdf:
        _hoja_a_pdf(pdf, hoja, fotos_hoja, recursos_pdf)
    pdf.save()
    for hoja in libro.worksheets:
        _quitar_metaarchivos(hoja)
    excel_buffer = io.BytesIO()
    libro.save(excel_buffer)
    resumen = {
        "hojas": hojas,
        "actividades": len(actividades),
        "fotos": fotos_insertadas,
        "omitidas": fotos_omitidas,
    }
    return excel_buffer.getvalue(), pdf_buffer.getvalue(), resumen


def cajas_desde_libro(contenido: bytes) -> list[dict]:
    """Lee las fotos de evidencia ancladas en la cuadrícula 3xN."""
    with warnings.catch_warnings():
        warnings.filterwarnings("ignore", category=UserWarning)
        libro = load_workbook(io.BytesIO(contenido))
    cajas = []
    for hoja in libro.worksheets:
        for dibujo in hoja._images:
            ancla = dibujo.anchor
            if not isinstance(ancla, OneCellAnchor):
                continue
            fila = ancla._from.row + 1
            if fila < 10 or fila > 53:
                continue
            columna = ancla._from.col + 1
            dx = ancla._from.colOff / 9525
            dy = ancla._from.rowOff / 9525
            ancho = ancla.ext.cx / 9525
            alto = ancla.ext.cy / 9525
            origen_x, origen_y = _origen_px(hoja, columna, fila)
            ranura = next(
                item
                for item in RANURAS
                if item[0] == columna and item[1] == fila
            )
            caja_x, caja_y = _origen_px(hoja, ranura[0], ranura[1])
            caja_ancho, caja_alto = _tamano_rango_px(hoja, *ranura)
            cajas.append(
                {
                    "hoja": hoja.title,
                    "x": origen_x + dx,
                    "y": origen_y + dy,
                    "ancho": ancho,
                    "alto": alto,
                    "ranura": (caja_x, caja_y, caja_ancho, caja_alto),
                    "pixeles": (dibujo.width, dibujo.height),
                }
            )
    return cajas
