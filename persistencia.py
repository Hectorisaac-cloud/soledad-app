"""Respaldo permanente de la bitácora y limpieza de archivos pesados.

Streamlit Cloud reconstruye el disco con el contenido de `main` en cada
despliegue. Fechas, pozos, actividades, insumos y estatus se conservan en
la rama Git `datos`, que no dispara ese despliegue. Las fotografías y los
PDF de OT no se suben ni se vuelven a bajar: viven en el disco solo hasta
que el usuario descarga su reporte, y después se retiran.
"""

from __future__ import annotations

import base64
import errno
import hashlib
import io
import json
import os
import shutil
import stat
import tempfile
import threading
import urllib.error
import urllib.request
import zipfile
from datetime import datetime

RAIZ = os.path.dirname(os.path.abspath(__file__))
RAMA_DATOS_DEFAULT = "datos"
REPO_DEFAULT = "Hectorisaac-cloud/soledad-app"
ARCHIVO_BITACORA = "database/bitacora_operativa_soledad.xlsx"
ARCHIVO_ESTADO = "database/persistencia_estado.json"
ARCHIVO_EVIDENCIAS_NUEVAS = "database/evidencias_pendientes.txt"
ARCHIVO_EVIDENCIAS_ELIMINADAS = "database/evidencias_eliminadas.txt"
ARCHIVO_BAJAS_BITACORA = "database/bitacora_bajas.json"
ARCHIVO_OT_LIBERADAS = "database/ot_liberadas.json"
ARCHIVOS_DATOS = (
    ARCHIVO_BITACORA,
    "database/programa_norte.xlsx",
    "database/programa_sur.xlsx",
    "database/programa_mantenimiento.xlsx",
    ARCHIVO_EVIDENCIAS_ELIMINADAS,
    "baul_evidencias.json",
    "evidencia_fotografica.json",
    "almacen_insumos.json",
    "asignaciones_diarias.json",
    "catalogos.json",
    "ubicaciones_equipos.json",
    "control_programa_octubre.json",
    "programados_ops.json",
    "actividades_cuadrillas.json",
    "personal_cuadrillas.json",
    "vehiculos_ot.json",
    "soporte_refacciones.json",
    ARCHIVO_OT_LIBERADAS,
)
CARPETAS_EVIDENCIA = (
    "evidencia_fotografica_octubre",
    "Baul_Evidencias",
    "evidencias",
    "Evidencia Fotografica",
    os.path.join("taller_vehicular", "evidencias"),
)
CARPETAS_PDF = (
    "bitacora_pdf",
    "historial_ot",
)
CARPETAS_PESADAS = CARPETAS_EVIDENCIA + CARPETAS_PDF
EXTENSIONES_EVIDENCIA = {".jpg", ".jpeg", ".png", ".webp", ".gif", ".bmp"}
EXTENSIONES_PESADAS = EXTENSIONES_EVIDENCIA | {".pdf"}
MAX_ARCHIVOS_POR_RESPALDO = 25
MAX_BYTES_POR_RESPALDO = 20 * 1024 * 1024
MAX_BYTES_ZIP = 25 * 1024 * 1024
MAX_PURGA_PESADOS = 200
API = "https://api.github.com"
_CANDADO = threading.Lock()


def elegir_bitacora(filas_locales: int, filas_remotas: int, bajas: int = 0) -> str:
    """La copia con más registros gana. Un empate conserva la local.

    `bajas` son los registros que se eliminaron a propósito en esta copia y
    todavía no llegan a la rama de respaldo; no cuentan como pérdida.
    """
    if filas_remotas > filas_locales + max(0, int(bajas or 0)):
        return "remota"
    return "local"


def leer_bajas_bitacora() -> int:
    ruta = resolver_ruta(ARCHIVO_BAJAS_BITACORA)
    if not os.path.exists(ruta):
        return 0
    try:
        with open(ruta, encoding="utf-8") as archivo:
            return max(0, int((json.load(archivo) or {}).get("pendientes", 0)))
    except (OSError, ValueError, TypeError, AttributeError):
        return 0


def _escribir_bajas_bitacora(pendientes: int) -> None:
    destino = resolver_ruta(ARCHIVO_BAJAS_BITACORA)
    os.makedirs(os.path.dirname(destino), exist_ok=True)
    contenido = json.dumps(
        {"pendientes": max(0, int(pendientes)), "cuando": datetime.now().isoformat(timespec="seconds")},
        ensure_ascii=False,
    ).encode("utf-8")
    _escribir_atomico(destino, contenido)


def registrar_baja_bitacora(cantidad: int = 1) -> None:
    _escribir_bajas_bitacora(leer_bajas_bitacora() + max(0, int(cantidad)))


def _evidencias_eliminadas() -> set[str]:
    ruta = resolver_ruta(ARCHIVO_EVIDENCIAS_ELIMINADAS)
    if not os.path.exists(ruta):
        return set()
    with open(ruta, encoding="utf-8") as archivo:
        return {linea.strip().replace("\\", "/") for linea in archivo if linea.strip()}


def _prefijo_pesado(ruta: str) -> bool:
    normal = str(ruta or "").replace("\\", "/")
    for carpeta in CARPETAS_PESADAS:
        base = carpeta.replace("\\", "/").rstrip("/")
        if normal == base or normal.startswith(base + "/"):
            return True
    return False


def _prefijo_evidencia(ruta: str) -> bool:
    return _prefijo_pesado(ruta)


def _relativa_pesada(ruta: str) -> str | None:
    """Ruta de foto o PDF dentro del proyecto. Nada fuera de la app."""
    texto = str(ruta or "").replace("\\", "/").strip()
    if not texto:
        return None
    if os.path.isabs(texto) or (len(texto) > 2 and texto[1] == ":"):
        try:
            texto = os.path.relpath(texto, RAIZ).replace("\\", "/")
        except ValueError:
            return None
    segura = ruta_segura(texto)
    if not segura or not _prefijo_pesado(segura):
        return None
    if os.path.splitext(segura)[1].lower() not in EXTENSIONES_PESADAS:
        return None
    absoluto = os.path.abspath(os.path.join(RAIZ, segura.replace("/", os.sep)))
    try:
        if os.path.commonpath([os.path.abspath(RAIZ), absoluto]) != os.path.abspath(RAIZ):
            return None
    except ValueError:
        return None
    return segura


def es_archivo_pesado(ruta: str) -> bool:
    return _relativa_pesada(ruta) is not None


def _relativa_evidencia(ruta: str) -> str | None:
    return _relativa_pesada(ruta)


def _reescribir_lista(relativa: str, lineas: list[str]) -> None:
    destino = resolver_ruta(relativa)
    os.makedirs(os.path.dirname(destino), exist_ok=True)
    texto = "".join(f"{linea}\n" for linea in lineas)
    _escribir_atomico(destino, texto.encode("utf-8"))


def marcar_evidencia_eliminada(rutas: list[str] | None) -> None:
    """Anota las fotos borradas para que el respaldo no las vuelva a descargar."""
    limpias = [rel for rel in (_relativa_evidencia(r) for r in rutas or []) if rel]
    if not limpias:
        return
    actuales = sorted(_evidencias_eliminadas() | set(limpias))
    _reescribir_lista(ARCHIVO_EVIDENCIAS_ELIMINADAS, actuales)
    pendientes = _evidencias_marcadas()
    if pendientes:
        quitar = set(limpias)
        _reescribir_lista(ARCHIVO_EVIDENCIAS_NUEVAS, [p for p in pendientes if p not in quitar])


def _sin_solo_lectura(ruta: str) -> None:
    try:
        modo = os.stat(ruta).st_mode
        if not modo & stat.S_IWRITE:
            os.chmod(ruta, modo | stat.S_IWRITE)
    except OSError:
        pass


def borrar_archivo(ruta: str) -> None:
    """os.remove que también borra archivos marcados como solo lectura (OneDrive)."""
    try:
        os.remove(ruta)
    except PermissionError:
        _sin_solo_lectura(ruta)
        os.remove(ruta)


def borrar_carpeta_vacia(ruta: str) -> None:
    """os.rmdir que también borra carpetas con atributo de solo lectura (OneDrive)."""
    try:
        os.rmdir(ruta)
    except PermissionError:
        _sin_solo_lectura(ruta)
        os.rmdir(ruta)


def aplicar_evidencias_eliminadas() -> int:
    """Borra del disco las fotos que se eliminaron en una sesión anterior."""
    borradas = 0
    for relativa in _evidencias_eliminadas():
        destino = resolver_ruta(relativa)
        if os.path.isfile(destino):
            try:
                borrar_archivo(destino)
                borradas += 1
            except OSError:
                continue
    return borradas


def debe_bloquear_vaciado(filas_actuales: int, filas_nuevas: int) -> bool:
    if filas_nuevas > 0:
        return False
    return filas_actuales != 0


def git_blob_sha(contenido: bytes) -> str:
    encabezado = f"blob {len(contenido)}\0".encode("ascii")
    return hashlib.sha1(encabezado + contenido).hexdigest()


def ruta_segura(relativa: str) -> str | None:
    texto = str(relativa or "").replace("\\", "/").strip()
    if texto.startswith("./"):
        texto = texto[2:]
    texto = texto.lstrip("/")
    if not texto or texto.startswith("../"):
        return None
    partes = [parte for parte in texto.split("/") if parte and parte != "."]
    if any(parte == ".." for parte in partes):
        return None
    return "/".join(partes)


def resolver_ruta(relativa: str) -> str:
    segura = ruta_segura(relativa) or ""
    directa = segura.replace("/", os.sep)
    if directa and os.path.exists(directa):
        return directa
    return os.path.join(RAIZ, directa) if directa else os.path.join(RAIZ, ARCHIVO_BITACORA)


def _secreto(nombre: str) -> str:
    entorno = os.environ.get(nombre, "").strip()
    if entorno:
        return entorno
    try:
        import streamlit as st

        valor = st.secrets.get(nombre, "")
    except Exception:
        return ""
    return str(valor or "").strip()


def credenciales_github() -> dict | None:
    token = _secreto("GITHUB_TOKEN") or _secreto("GH_TOKEN")
    if not token:
        return None
    repo = _secreto("GITHUB_REPO") or REPO_DEFAULT
    rama = _secreto("GITHUB_RAMA_DATOS") or RAMA_DATOS_DEFAULT
    if "/" not in repo or not rama or rama in {"main", "master"}:
        return None
    return {"token": token, "repo": repo, "rama": rama}


def contar_filas_bitacora(ruta: str | None) -> int:
    if not ruta or not os.path.exists(ruta) or os.path.getsize(ruta) <= 0:
        return 0
    try:
        from openpyxl import load_workbook
    except Exception:
        return 0
    wb = load_workbook(ruta, read_only=True, data_only=True)
    try:
        nombres = {nombre.casefold(): nombre for nombre in wb.sheetnames}
        hoja = None
        for preferida in ("bitacora", "hoja5", "hoja2"):
            if preferida in nombres:
                hoja = nombres[preferida]
                break
        ws = wb[hoja or wb.sheetnames[0]]
        col_fecha = None
        fila_header = None
        total = 0
        for indice, fila in enumerate(ws.iter_rows(values_only=True)):
            if col_fecha is None and indice < 12:
                for columna, valor in enumerate(fila):
                    if str(valor or "").strip().casefold() == "fecha":
                        col_fecha = columna
                        fila_header = indice
                        break
                continue
            if col_fecha is None or fila_header is None or indice <= fila_header:
                continue
            if col_fecha < len(fila) and fila[col_fecha] not in (None, ""):
                total += 1
        return total
    except Exception:
        return -1
    finally:
        wb.close()


def _es_cruce_de_dispositivo(exc: BaseException) -> bool:
    return isinstance(exc, OSError) and getattr(exc, "errno", None) in {errno.EXDEV, 18}


def copiar_seguro(origen: str, destino: str) -> None:
    """Copia bytes aunque origen y destino estén en discos distintos."""
    carpeta = os.path.dirname(os.path.abspath(destino))
    if carpeta:
        os.makedirs(carpeta, exist_ok=True)
    try:
        shutil.copy(origen, destino)
        return
    except OSError as exc:
        if not _es_cruce_de_dispositivo(exc):
            raise
    with open(origen, "rb") as entrada, open(destino, "wb") as salida:
        shutil.copyfileobj(entrada, salida)


def reemplazar_seguro(origen: str, destino: str) -> None:
    """Sustituye el destino. Si el renombrado cruza de partición, copia y borra el temporal."""
    try:
        os.replace(origen, destino)
        return
    except OSError as exc:
        if not _es_cruce_de_dispositivo(exc):
            raise
    copiar_seguro(origen, destino)
    try:
        os.remove(origen)
    except OSError:
        pass


def _guardar_estado(estado: dict) -> None:
    destino = resolver_ruta(ARCHIVO_ESTADO)
    os.makedirs(os.path.dirname(destino), exist_ok=True)
    temporal = destino + ".tmp"
    with open(temporal, "w", encoding="utf-8") as archivo:
        json.dump(estado, archivo, ensure_ascii=False, indent=2)
    reemplazar_seguro(temporal, destino)


def leer_estado() -> dict:
    ruta = resolver_ruta(ARCHIVO_ESTADO)
    if not os.path.exists(ruta):
        return {}
    try:
        with open(ruta, encoding="utf-8") as archivo:
            datos = json.load(archivo)
    except Exception:
        return {}
    return datos if isinstance(datos, dict) else {}


def _api(metodo: str, url: str, token: str, payload: dict | None = None, crudo: bool = False):
    datos = None if payload is None else json.dumps(payload).encode("utf-8")
    peticion = urllib.request.Request(url, data=datos, method=metodo)
    peticion.add_header("Authorization", f"Bearer {token}")
    peticion.add_header("User-Agent", "soledad-app-persistencia")
    peticion.add_header("X-GitHub-Api-Version", "2022-11-28")
    peticion.add_header("Accept", "application/vnd.github.raw" if crudo else "application/vnd.github+json")
    if payload is not None:
        peticion.add_header("Content-Type", "application/json")
    try:
        with urllib.request.urlopen(peticion, timeout=45) as respuesta:
            cuerpo = respuesta.read()
    except urllib.error.HTTPError as exc:
        detalle = exc.read().decode("utf-8", errors="replace")[:500]
        raise RuntimeError(f"GitHub respondió {exc.code}: {detalle}") from exc
    if crudo:
        return cuerpo
    if not cuerpo:
        return {}
    return json.loads(cuerpo.decode("utf-8"))


def _obtener_ref(cred: dict, rama: str) -> dict | None:
    url = f"{API}/repos/{cred['repo']}/git/ref/heads/{rama}"
    try:
        return _api("GET", url, cred["token"])
    except RuntimeError as exc:
        if "404" in str(exc):
            return None
        raise


def _asegurar_rama(cred: dict) -> str:
    ref = _obtener_ref(cred, cred["rama"])
    if ref:
        return ref["object"]["sha"]
    base = _obtener_ref(cred, "main") or _obtener_ref(cred, "master")
    if not base:
        raise RuntimeError("No se encontró la rama main para crear el respaldo.")
    creado = _api(
        "POST",
        f"{API}/repos/{cred['repo']}/git/refs",
        cred["token"],
        {"ref": f"refs/heads/{cred['rama']}", "sha": base["object"]["sha"]},
    )
    return creado["object"]["sha"]


def _arbol_remoto(cred: dict, commit_sha: str) -> dict[str, str]:
    commit = _api("GET", f"{API}/repos/{cred['repo']}/git/commits/{commit_sha}", cred["token"])
    arbol = _api(
        "GET",
        f"{API}/repos/{cred['repo']}/git/trees/{commit['tree']['sha']}?recursive=1",
        cred["token"],
    )
    mapa = {}
    for item in arbol.get("tree") or []:
        if item.get("type") == "blob" and item.get("path") and item.get("sha"):
            mapa[item["path"]] = item["sha"]
    return mapa


def _descargar_blob(cred: dict, sha: str) -> bytes:
    return _api("GET", f"{API}/repos/{cred['repo']}/git/blobs/{sha}", cred["token"], crudo=True)


def _directorio_temporal(destino: str | None = None) -> str:
    carpeta = os.path.dirname(destino) if destino else ""
    if carpeta:
        os.makedirs(carpeta, exist_ok=True)
        return carpeta
    return tempfile.gettempdir()


def _escribir_atomico(destino: str, contenido: bytes) -> None:
    os.makedirs(os.path.dirname(destino) or ".", exist_ok=True)
    fd, temporal = tempfile.mkstemp(suffix=".tmp", dir=_directorio_temporal(destino))
    os.close(fd)
    try:
        with open(temporal, "wb") as archivo:
            archivo.write(contenido)
        reemplazar_seguro(temporal, destino)
    finally:
        if os.path.exists(temporal):
            try:
                os.remove(temporal)
            except OSError:
                pass


def leer_ot_liberadas() -> dict[str, dict]:
    ruta = resolver_ruta(ARCHIVO_OT_LIBERADAS)
    if not os.path.exists(ruta):
        return {}
    try:
        with open(ruta, encoding="utf-8") as archivo:
            datos = json.load(archivo)
    except (OSError, ValueError, TypeError):
        return {}
    if not isinstance(datos, dict):
        return {}
    salida = {}
    for clave, valor in datos.items():
        fecha = str(clave or "").strip()
        if not fecha:
            continue
        if isinstance(valor, dict):
            salida[fecha] = {
                "nombre": str(valor.get("nombre") or "").strip(),
                "cuando": str(valor.get("cuando") or "").strip(),
            }
        else:
            salida[fecha] = {"nombre": str(valor or "").strip(), "cuando": ""}
    return salida


def _guardar_ot_liberadas(datos: dict[str, dict]) -> None:
    destino = resolver_ruta(ARCHIVO_OT_LIBERADAS)
    os.makedirs(os.path.dirname(destino), exist_ok=True)
    contenido = json.dumps(datos, ensure_ascii=False, indent=2).encode("utf-8")
    _escribir_atomico(destino, contenido)


def registrar_ot_liberada(fecha_iso: str, nombre: str) -> None:
    fecha = str(fecha_iso or "").strip()
    if not fecha:
        return
    datos = leer_ot_liberadas()
    datos[fecha] = {
        "nombre": str(nombre or "").strip(),
        "cuando": datetime.now().isoformat(timespec="seconds"),
    }
    _guardar_ot_liberadas(datos)


def olvidar_ot_liberada(fecha_iso: str) -> None:
    fecha = str(fecha_iso or "").strip()
    datos = leer_ot_liberadas()
    if fecha not in datos:
        return
    datos.pop(fecha, None)
    _guardar_ot_liberadas(datos)


def _podar_carpetas_vacias(relativa: str) -> None:
    actual = os.path.dirname(resolver_ruta(relativa))
    raiz = os.path.abspath(RAIZ)
    topes = {os.path.abspath(resolver_ruta(carpeta)) for carpeta in CARPETAS_PESADAS}
    topes.add(raiz)
    while actual:
        absoluto = os.path.abspath(actual)
        if absoluto in topes:
            break
        try:
            if os.path.commonpath([raiz, absoluto]) != raiz or os.listdir(absoluto):
                break
            borrar_carpeta_vacia(absoluto)
        except (OSError, ValueError):
            break
        actual = os.path.dirname(absoluto)


def liberar_archivos_pesados(rutas: list[str] | None) -> int:
    """Borra fotos y PDF ya descargados. No toca la bitácora ni los JSON."""
    limpias = []
    vistos = set()
    for ruta in rutas or []:
        relativa = _relativa_pesada(ruta)
        if not relativa or relativa in vistos:
            continue
        vistos.add(relativa)
        limpias.append(relativa)
    if not limpias:
        return 0
    borradas = 0
    for relativa in limpias:
        destino = resolver_ruta(relativa)
        if os.path.isfile(destino):
            try:
                borrar_archivo(destino)
                borradas += 1
            except OSError:
                continue
        _podar_carpetas_vacias(relativa)
    marcar_evidencia_eliminada(limpias)
    return borradas


def cambios_purga_pesados(arbol: dict[str, str] | None) -> list[dict]:
    """Quita fotos y PDF de la rama de respaldo sin borrar los registros de texto."""
    cambios = []
    for ruta in sorted(arbol or {}):
        if not es_archivo_pesado(ruta):
            continue
        cambios.append({"path": ruta, "mode": "100644", "type": "blob", "sha": None})
        if len(cambios) >= MAX_PURGA_PESADOS:
            break
    return cambios


def restaurar_desde_github() -> dict:
    cred = credenciales_github()
    if not cred:
        return {"ok": False, "motivo": "sin_token", "mensaje": "No hay GITHUB_TOKEN configurado."}
    if not _CANDADO.acquire(blocking=False):
        return {"ok": False, "motivo": "ocupado", "mensaje": "Ya hay una sincronización en curso."}
    try:
        ref = _obtener_ref(cred, cred["rama"])
        if not ref:
            estado = {
                "ok": True,
                "cuando": datetime.now().isoformat(timespec="seconds"),
                "mensaje": f"La rama {cred['rama']} todavía no existe. Se creará con el próximo registro.",
                "filas_locales": contar_filas_bitacora(resolver_ruta(ARCHIVO_BITACORA)),
            }
            _guardar_estado(estado)
            return estado
        arbol = _arbol_remoto(cred, ref["object"]["sha"])
        locales = contar_filas_bitacora(resolver_ruta(ARCHIVO_BITACORA))
        remotas = 0
        adopto_remota = False
        blob_bitacora = arbol.get(ARCHIVO_BITACORA)
        if blob_bitacora:
            contenido = _descargar_blob(cred, blob_bitacora)
            fd, temporal = tempfile.mkstemp(suffix=".xlsx", dir=_directorio_temporal(resolver_ruta(ARCHIVO_BITACORA)))
            os.close(fd)
            try:
                with open(temporal, "wb") as archivo:
                    archivo.write(contenido)
                remotas = contar_filas_bitacora(temporal)
                if locales >= 0 and remotas >= 0 and elegir_bitacora(locales, remotas, leer_bajas_bitacora()) == "remota":
                    _escribir_atomico(resolver_ruta(ARCHIVO_BITACORA), contenido)
                    locales = remotas
                    adopto_remota = True
            finally:
                if os.path.exists(temporal):
                    os.remove(temporal)
        recuperados = 0
        if adopto_remota:
            for relativa in ARCHIVOS_DATOS:
                if relativa == ARCHIVO_BITACORA or relativa not in arbol:
                    continue
                _escribir_atomico(resolver_ruta(relativa), _descargar_blob(cred, arbol[relativa]))
                recuperados += 1
        else:
            for relativa in ARCHIVOS_DATOS:
                if relativa == ARCHIVO_BITACORA or relativa not in arbol:
                    continue
                destino = resolver_ruta(relativa)
                if os.path.exists(destino) and os.path.getsize(destino) > 0:
                    continue
                _escribir_atomico(destino, _descargar_blob(cred, arbol[relativa]))
                recuperados += 1
        aplicar_evidencias_eliminadas()
        mensaje = (
            f"Bitácora local: {locales} registros. "
            f"Respaldo en {cred['rama']}: {remotas} registros. "
            f"Se recuperaron {recuperados} archivo(s) de datos. "
            "Las fotografías y los PDF no se vuelven a descargar."
        )
        estado = {
            "ok": True,
            "cuando": datetime.now().isoformat(timespec="seconds"),
            "mensaje": mensaje,
            "filas_locales": locales,
            "filas_remotas": remotas,
            "recuperados": recuperados,
            "pendientes": 0,
        }
        _guardar_estado(estado)
        return estado
    except Exception as exc:
        estado = {
            "ok": False,
            "cuando": datetime.now().isoformat(timespec="seconds"),
            "mensaje": f"No se pudo recuperar el respaldo: {exc}",
        }
        try:
            _guardar_estado(estado)
        except Exception:
            pass
        return estado
    finally:
        _CANDADO.release()


def restaurar_al_arrancar() -> dict:
    try:
        import streamlit as st

        if st.session_state.get("_persistencia_restaurada"):
            return {"ok": True, "omitido": True}
        st.session_state["_persistencia_restaurada"] = True
    except Exception:
        pass
    estado = restaurar_desde_github()
    try:
        aplicar_evidencias_eliminadas()
    except Exception:
        pass
    return estado


def marcar_evidencia_nueva(rutas: list[str] | None) -> None:
    limpias = []
    for ruta in rutas or []:
        relativa = _relativa_pesada(ruta)
        if relativa:
            limpias.append(relativa)
    if not limpias:
        return
    eliminadas = _evidencias_eliminadas()
    if eliminadas & set(limpias):
        _reescribir_lista(ARCHIVO_EVIDENCIAS_ELIMINADAS, sorted(eliminadas - set(limpias)))
    destino = resolver_ruta(ARCHIVO_EVIDENCIAS_NUEVAS)
    os.makedirs(os.path.dirname(destino), exist_ok=True)
    previas = []
    if os.path.exists(destino):
        with open(destino, encoding="utf-8") as archivo:
            previas = [linea.strip() for linea in archivo if linea.strip()]
    vistos = set(previas)
    with open(destino, "a", encoding="utf-8") as archivo:
        for ruta in limpias:
            if ruta not in vistos:
                archivo.write(ruta + "\n")
                vistos.add(ruta)


def _evidencias_marcadas() -> list[str]:
    ruta = resolver_ruta(ARCHIVO_EVIDENCIAS_NUEVAS)
    if not os.path.exists(ruta):
        return []
    with open(ruta, encoding="utf-8") as archivo:
        return [linea.strip().replace("\\", "/") for linea in archivo if linea.strip()]


def _pesados_locales() -> list[str]:
    encontrados = []
    vistos = set()
    for carpeta in CARPETAS_PESADAS:
        base = resolver_ruta(carpeta)
        if not os.path.isdir(base):
            continue
        for raiz, _, nombres in os.walk(base):
            for nombre in nombres:
                if nombre.startswith("~$") or nombre.startswith("."):
                    continue
                absoluto = os.path.join(raiz, nombre)
                relativa = _relativa_pesada(os.path.relpath(absoluto, RAIZ))
                if not relativa or relativa in vistos or not os.path.isfile(resolver_ruta(relativa)):
                    continue
                vistos.add(relativa)
                encontrados.append(relativa)
    return encontrados


def _archivos_a_respaldar(arbol: dict[str, str] | None = None, *, solo_copia_local: bool = False) -> list[str]:
    """La rama de respaldo solo recibe datos de texto. El ZIP local sí puede llevar fotos y PDF."""
    del arbol
    elegidos = []
    for relativa in ARCHIVOS_DATOS:
        if os.path.exists(resolver_ruta(relativa)):
            elegidos.append(relativa.replace("\\", "/"))
    if solo_copia_local:
        for relativa in _pesados_locales():
            if relativa not in elegidos:
                elegidos.append(relativa)
    return elegidos


def respaldar_cambios(motivo: str) -> dict:
    cred = credenciales_github()
    if not cred:
        estado = {
            "ok": False,
            "motivo": "sin_token",
            "cuando": datetime.now().isoformat(timespec="seconds"),
            "mensaje": "El registro quedó en el servidor, sin copia en GitHub. Configura GITHUB_TOKEN para que sobreviva a un redespliegue.",
            "filas_locales": contar_filas_bitacora(resolver_ruta(ARCHIVO_BITACORA)),
        }
        try:
            _guardar_estado(estado)
        except Exception:
            pass
        return estado
    if not _CANDADO.acquire(blocking=False):
        return {"ok": False, "motivo": "ocupado", "mensaje": "Ya hay una sincronización en curso."}
    try:
        commit_sha = _asegurar_rama(cred)
        arbol = _arbol_remoto(cred, commit_sha)
        locales = contar_filas_bitacora(resolver_ruta(ARCHIVO_BITACORA))
        blob_remoto = arbol.get(ARCHIVO_BITACORA)
        if blob_remoto:
            fd, temporal = tempfile.mkstemp(suffix=".xlsx", dir=_directorio_temporal(resolver_ruta(ARCHIVO_BITACORA)))
            os.close(fd)
            try:
                with open(temporal, "wb") as archivo:
                    archivo.write(_descargar_blob(cred, blob_remoto))
                remotas = contar_filas_bitacora(temporal)
            finally:
                if os.path.exists(temporal):
                    os.remove(temporal)
            if locales < 0 or remotas < 0:
                estado = {
                    "ok": False,
                    "motivo": "conteo",
                    "cuando": datetime.now().isoformat(timespec="seconds"),
                    "mensaje": "No se subió la bitácora porque no se pudo contar una de las dos copias.",
                    "filas_locales": locales,
                    "filas_remotas": remotas,
                }
                _guardar_estado(estado)
                return estado
            if remotas > locales + leer_bajas_bitacora():
                estado = {
                    "ok": False,
                    "motivo": "remoto_mas_nuevo",
                    "cuando": datetime.now().isoformat(timespec="seconds"),
                    "mensaje": (
                        f"No se subió la bitácora: la rama {cred['rama']} tiene {remotas} registros "
                        f"y esta copia tiene {locales}. Recupera el respaldo antes de seguir capturando."
                    ),
                    "filas_locales": locales,
                    "filas_remotas": remotas,
                }
                _guardar_estado(estado)
                return estado
        cambios = []
        bytes_usados = 0
        omitidos = 0
        datos_permanentes = {item.replace("\\", "/") for item in ARCHIVOS_DATOS}
        for relativa in _archivos_a_respaldar(arbol):
            if es_archivo_pesado(relativa):
                continue
            if relativa == ARCHIVO_BITACORA and locales <= 0:
                continue
            absoluto = resolver_ruta(relativa)
            try:
                with open(absoluto, "rb") as archivo:
                    contenido = archivo.read()
            except OSError:
                continue
            if not contenido:
                continue
            firma = git_blob_sha(contenido)
            if arbol.get(relativa) == firma:
                continue
            es_dato = relativa in datos_permanentes
            if not es_dato and (
                len(cambios) >= MAX_ARCHIVOS_POR_RESPALDO or bytes_usados + len(contenido) > MAX_BYTES_POR_RESPALDO
            ):
                omitidos += 1
                continue
            blob = _api(
                "POST",
                f"{API}/repos/{cred['repo']}/git/blobs",
                cred["token"],
                {"content": base64.b64encode(contenido).decode("ascii"), "encoding": "base64"},
            )
            cambios.append(
                {"path": relativa, "mode": "100644", "type": "blob", "sha": blob["sha"]}
            )
            bytes_usados += len(contenido)
        ya = {item["path"] for item in cambios}
        for relativa in sorted(_evidencias_eliminadas()):
            if relativa in arbol and relativa not in ya and not os.path.exists(resolver_ruta(relativa)):
                cambios.append({"path": relativa, "mode": "100644", "type": "blob", "sha": None})
                ya.add(relativa)
        for item in cambios_purga_pesados(arbol):
            if item["path"] not in ya:
                cambios.append(item)
                ya.add(item["path"])
        if not cambios:
            estado = {
                "ok": True,
                "cuando": datetime.now().isoformat(timespec="seconds"),
                "mensaje": f"La rama {cred['rama']} ya tiene esta bitácora ({locales} registros).",
                "filas_locales": locales,
                "subidos": 0,
            }
            _guardar_estado(estado)
            return estado
        commit = _api("GET", f"{API}/repos/{cred['repo']}/git/commits/{commit_sha}", cred["token"])
        arbol_nuevo = _api(
            "POST",
            f"{API}/repos/{cred['repo']}/git/trees",
            cred["token"],
            {"base_tree": commit["tree"]["sha"], "tree": cambios},
        )
        nuevo = _api(
            "POST",
            f"{API}/repos/{cred['repo']}/git/commits",
            cred["token"],
            {
                "message": f"Respaldo operativo Soledad\n\n{motivo}. Registros: {locales}.",
                "tree": arbol_nuevo["sha"],
                "parents": [commit_sha],
            },
        )
        _api(
            "PATCH",
            f"{API}/repos/{cred['repo']}/git/refs/heads/{cred['rama']}",
            cred["token"],
            {"sha": nuevo["sha"]},
        )
        if not omitidos and leer_bajas_bitacora():
            _escribir_bajas_bitacora(0)
        purgados = sum(1 for item in cambios if item.get("sha") is None)
        mensaje = (
            f"Respaldo enviado a la rama {cred['rama']}: {locales} registros de bitácora "
            f"y {len(cambios) - purgados} archivo(s) de datos."
        )
        if purgados:
            mensaje += f" Se retiraron {purgados} foto(s) o PDF del respaldo."
        if omitidos:
            mensaje += f" Quedaron {omitidos} archivo(s) para el siguiente respaldo."
        estado = {
            "ok": True,
            "cuando": datetime.now().isoformat(timespec="seconds"),
            "mensaje": mensaje,
            "filas_locales": locales,
            "subidos": len(cambios),
            "omitidos": omitidos,
            "commit": nuevo.get("sha", ""),
        }
        _guardar_estado(estado)
        return estado
    except Exception as exc:
        estado = {
            "ok": False,
            "cuando": datetime.now().isoformat(timespec="seconds"),
            "mensaje": f"No se pudo respaldar en GitHub: {exc}",
            "filas_locales": contar_filas_bitacora(resolver_ruta(ARCHIVO_BITACORA)),
        }
        try:
            _guardar_estado(estado)
        except Exception:
            pass
        return estado
    finally:
        _CANDADO.release()


def _leer_bytes(absoluto: str) -> bytes:
    """Lee un archivo. Si el disco rechaza el acceso directo, lo copia a un temporal."""
    try:
        with open(absoluto, "rb") as archivo:
            return archivo.read()
    except OSError as exc:
        if not _es_cruce_de_dispositivo(exc):
            raise
    fd, temporal = tempfile.mkstemp(prefix="soledad_lectura_")
    os.close(fd)
    try:
        copiar_seguro(absoluto, temporal)
        with open(temporal, "rb") as archivo:
            return archivo.read()
    finally:
        if os.path.exists(temporal):
            try:
                os.remove(temporal)
            except OSError:
                pass


def construir_copia_seguridad() -> tuple[bytes, str, list[str]]:
    sello = datetime.now().strftime("%Y%m%d_%H%M%S")
    nombre = f"respaldo_bitacora_{sello}.zip"
    buffer = io.BytesIO()
    incluidos = []
    pesados = []
    omitidos = []
    usados = 0
    with zipfile.ZipFile(buffer, "w", compression=zipfile.ZIP_DEFLATED) as paquete:
        for relativa in _archivos_a_respaldar(solo_copia_local=True):
            absoluto = resolver_ruta(relativa)
            try:
                tamano = os.path.getsize(absoluto)
            except OSError:
                continue
            if usados + tamano > MAX_BYTES_ZIP:
                omitidos.append(relativa)
                continue
            arcname = relativa.replace("\\", "/")
            try:
                paquete.writestr(arcname, _leer_bytes(absoluto))
            except OSError:
                omitidos.append(relativa)
                continue
            incluidos.append(arcname)
            usados += tamano
            if es_archivo_pesado(relativa):
                pesados.append(arcname)
        nota = [
            "Copia de seguridad de Soledad",
            f"Generada: {datetime.now().isoformat(timespec='seconds')}",
            f"Registros en la bitácora: {contar_filas_bitacora(resolver_ruta(ARCHIVO_BITACORA))}",
            "",
            "La bitácora y los JSON (fechas, pozos, actividades, insumos y estatus)",
            "se conservan en la rama datos. Las fotografías y los PDF viajan en esta",
            "copia para descargarlos; al descargarla, la app los retira del servidor.",
            "No conectes la rama datos al despliegue de Streamlit.",
            "",
            "Archivos incluidos:",
            *incluidos,
        ]
        if omitidos:
            nota.extend(["", "No cupieron en esta copia y siguen en el servidor:", *omitidos[:200]])
        paquete.writestr("LEEME_RESPALDO.txt", "\n".join(nota))
    return buffer.getvalue(), nombre, pesados
