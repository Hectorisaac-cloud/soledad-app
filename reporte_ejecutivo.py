"""Reporte ejecutivo para OPS a partir de la plantilla oficial."""

from __future__ import annotations

import io
import os
from datetime import date, datetime
from xml.sax.saxutils import escape

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.worksheet.worksheet import Worksheet
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import letter, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

RAIZ = os.path.dirname(os.path.abspath(__file__))
PLANTILLA_XLSX = os.path.join(RAIZ, "assets", "plantilla_reporte_ejecutivo.xlsx")
HOJA_PLANTILLA = "Hoja2"
FILA_ENCABEZADO = 4
FILA_DATOS = 5
COLUMNAS_EXCEL = ("C", "D", "E", "F", "G", "H", "I")
CAMPOS = ("ot", "fecha", "ubicacion", "cuadrilla", "descripcion", "insumo", "cantidad")
ENCABEZADOS = (
    "N° DE OT",
    "FECHA",
    "UBICACIÓN/POZO",
    "CUADRILLA/PERSONAL ASIGNADO",
    "DESCRIPCIÓN DE LA ACTIVIDAD",
    "INSUMO/REFACCIÓN",
    "CANTIDAD",
)


def _titulo(criterio: str) -> str:
    return f"REPORTE DE {str(criterio or '').strip().upper()}"


def _texto(valor) -> str:
    if valor is None:
        return ""
    if isinstance(valor, datetime):
        return valor.strftime("%d/%m/%Y")
    if isinstance(valor, date):
        return valor.strftime("%d/%m/%Y")
    return str(valor).strip()


def _cantidad(valor):
    if valor is None or valor == "":
        return None
    try:
        numero = float(valor)
    except (TypeError, ValueError):
        return None
    if numero <= 0:
        return None
    if abs(numero - round(numero)) < 1e-6:
        return int(round(numero))
    return round(numero, 2)


def _borde_formal() -> Border:
    linea = Side(style="thin", color="1F2937")
    return Border(left=linea, right=linea, top=linea, bottom=linea)


def _escribir_fila(ws: Worksheet, renglon: int, fila: dict, borde: Border) -> None:
    fuente = Font(name="Aptos Narrow", size=11, color="1F2937")
    alineacion = Alignment(horizontal="center", vertical="center", wrap_text=True)
    valores = (
        _texto(fila.get("ot")),
        fila.get("fecha") if isinstance(fila.get("fecha"), date) else _texto(fila.get("fecha")),
        _texto(fila.get("ubicacion")),
        _texto(fila.get("cuadrilla")),
        _texto(fila.get("descripcion")),
        _texto(fila.get("insumo")),
        _cantidad(fila.get("cantidad")),
    )
    for columna, valor in zip(COLUMNAS_EXCEL, valores):
        celda = ws[f"{columna}{renglon}"]
        celda.value = valor if valor is not None else None
        celda.font = fuente
        celda.alignment = alineacion
        celda.border = borde
        if columna == "D" and isinstance(valor, date):
            celda.number_format = "DD/MM/YYYY"
        if columna == "I" and isinstance(valor, (int, float)):
            celda.number_format = "0" if isinstance(valor, int) else "0.00"
    ws.row_dimensions[renglon].height = 30


def _hoja_destino(libro):
    if HOJA_PLANTILLA not in libro.sheetnames:
        raise KeyError(f"La plantilla no tiene la hoja «{HOJA_PLANTILLA}».")
    return libro[HOJA_PLANTILLA]


def _limpiar_referencias_externas(libro) -> None:
    """La plantilla arrastra nombres y vínculos de otros libros. Excel los refresca y deja la hoja en blanco."""
    libro._external_links = []
    for nombre in list(libro.defined_names):
        del libro.defined_names[nombre]


def _extender_tabla(hoja: Worksheet, ultima_fila: int) -> None:
    """La tabla oficial nace en C4:I5. Hay que crecerla para que Excel conserve las filas nuevas."""
    ref = f"C{FILA_ENCABEZADO}:I{max(ultima_fila, FILA_DATOS)}"
    if not hoja.tables:
        return
    for tabla in hoja.tables.values():
        tabla.ref = ref
        if tabla.autoFilter is not None:
            tabla.autoFilter.ref = ref


def generar_excel_reporte_ejecutivo(filas: list[dict], criterio: str, periodo: str) -> bytes:
    if not os.path.isfile(PLANTILLA_XLSX):
        raise FileNotFoundError("No está la plantilla assets/plantilla_reporte_ejecutivo.xlsx.")
    libro = load_workbook(PLANTILLA_XLSX)
    _limpiar_referencias_externas(libro)
    hoja = _hoja_destino(libro)
    borde = _borde_formal()
    hoja["C3"] = f"{_titulo(criterio)}\n{periodo}".strip()
    hoja["C3"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    hoja.row_dimensions[3].height = 36
    for columna in COLUMNAS_EXCEL:
        encabezado = hoja[f"{columna}{FILA_ENCABEZADO}"]
        encabezado.font = Font(name="Aptos Narrow", size=11, bold=True, color="1F2937")
        encabezado.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        encabezado.border = borde
    hoja.row_dimensions[FILA_ENCABEZADO].height = 32
    for indice, fila in enumerate(filas):
        _escribir_fila(hoja, FILA_DATOS + indice, fila, borde)
    ultima = FILA_DATOS + len(filas) - 1 if filas else FILA_DATOS
    _extender_tabla(hoja, ultima)
    hoja.print_area = f"C3:I{ultima}"
    hoja.print_title_rows = "3:4"
    hoja.page_setup.orientation = "landscape"
    hoja.page_setup.paperSize = hoja.PAPERSIZE_LETTER
    hoja.page_setup.fitToWidth = 1
    hoja.page_setup.fitToHeight = 1
    hoja.page_setup.horizontalCentered = True
    hoja.sheet_properties.pageSetUpPr.fitToPage = True
    hoja.page_margins.left = 0.4
    hoja.page_margins.right = 0.4
    hoja.page_margins.top = 0.5
    hoja.page_margins.bottom = 0.5
    hoja.oddFooter.center.text = "PJP4 - Area Contractual Soledad - OPS"
    hoja.sheet_view.topLeftCell = "A1"
    hoja.sheet_view.zoomScale = 110
    if hoja.sheet_view.selection:
        hoja.sheet_view.selection[0].activeCell = "C3"
        hoja.sheet_view.selection[0].sqref = "C3"
    buffer = io.BytesIO()
    libro.save(buffer)
    return buffer.getvalue()


def generar_pdf_reporte_ejecutivo(filas: list[dict], criterio: str, periodo: str) -> bytes:
    buffer = io.BytesIO()
    documento = SimpleDocTemplate(
        buffer,
        pagesize=landscape(letter),
        leftMargin=28,
        rightMargin=28,
        topMargin=26,
        bottomMargin=26,
        title=_titulo(criterio),
    )
    titulo = ParagraphStyle(
        "titulo_ejecutivo",
        fontName="Helvetica-Bold",
        fontSize=13,
        leading=16,
        alignment=TA_CENTER,
        textColor=colors.HexColor("#1F2937"),
    )
    subtitulo = ParagraphStyle(
        "subtitulo_ejecutivo",
        fontName="Helvetica",
        fontSize=9,
        leading=12,
        alignment=TA_CENTER,
        textColor=colors.HexColor("#374151"),
    )
    cabeza = ParagraphStyle(
        "cabeza_ejecutivo",
        fontName="Helvetica-Bold",
        fontSize=8,
        leading=10,
        alignment=TA_CENTER,
        textColor=colors.white,
    )
    cuerpo = ParagraphStyle(
        "cuerpo_ejecutivo",
        fontName="Helvetica",
        fontSize=8,
        leading=10,
        alignment=TA_CENTER,
        textColor=colors.HexColor("#1F2937"),
    )
    tabla_datos = [[Paragraph(escape(texto), cabeza) for texto in ENCABEZADOS]]
    if not filas:
        tabla_datos.append([Paragraph("Sin registros en el periodo seleccionado.", cuerpo)] + [""] * 6)
    for fila in filas:
        cantidad = _cantidad(fila.get("cantidad"))
        tabla_datos.append(
            [
                Paragraph(escape(_texto(fila.get("ot"))), cuerpo),
                Paragraph(escape(_texto(fila.get("fecha"))), cuerpo),
                Paragraph(escape(_texto(fila.get("ubicacion"))), cuerpo),
                Paragraph(escape(_texto(fila.get("cuadrilla"))), cuerpo),
                Paragraph(escape(_texto(fila.get("descripcion"))), cuerpo),
                Paragraph(escape(_texto(fila.get("insumo"))), cuerpo),
                Paragraph("" if cantidad is None else escape(str(cantidad)), cuerpo),
            ]
        )
    tabla = Table(
        tabla_datos,
        colWidths=[52, 68, 88, 120, 210, 110, 58],
        repeatRows=1,
    )
    estilo = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F2937")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#1F2937")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("LEFTPADDING", (0, 0), (-1, -1), 3),
        ("RIGHTPADDING", (0, 0), (-1, -1), 3),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
    ]
    if not filas:
        estilo.append(("SPAN", (0, 1), (-1, 1)))
    tabla.setStyle(TableStyle(estilo))
    documento.build(
        [
            Paragraph(escape(_titulo(criterio)), titulo),
            Spacer(1, 4),
            Paragraph(escape(f"{periodo} · PJP4 · Área Contractual Soledad"), subtitulo),
            Spacer(1, 10),
            tabla,
        ]
    )
    return buffer.getvalue()
